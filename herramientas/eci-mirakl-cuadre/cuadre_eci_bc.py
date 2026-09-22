# -*- coding: utf-8 -*-
"""
Cuadre de pedidos y devoluciones El Corte Inglés (marketplace Mirakl) vs Business Central.

Entradas
  --bc      'Mirakl Orders Import.xlsx': log del conector Mirakl → BC, una fila por línea de
            pedido ('Mirakl Line Id'), con 'Processing Status', 'BC Sales Order No.',
            'Return Status', 'Return Processing Status', 'BC Return Order No.', 'Error Message'…
  --mirakl  'pedidos 2026.xlsx': export de pedidos del portal de vendedor de ECI (Mirakl),
            hoja 'orders', una fila por línea ('N.º de asiento de pedido'), en hora de Madrid.
Salida
  --salida  Excel con las hojas Resumen, Cruce líneas, Cruce pedidos, Devoluciones,
            Export Mirakl (periodo, sin datos personales) y Log BC (periodo, sin nombres).
            Las columnas de cruce son fórmulas; los colores, formato condicional por Resultado.

Uso
  python cuadre_eci_bc.py --bc "Mirakl Orders Import.xlsx" --mirakl "pedidos 2026.xlsx" --periodo 2026-08
"""
import argparse
import calendar
import datetime as dt
import shutil
import sys
import warnings
import zipfile
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore")

FUENTE = "Arial"
MAXF = 6000  # filas máximas que cubren las fórmulas

COLOR = {
    "Coincide: pedido BC creado": "C6EFCE",
    "Coincide con diferencia de importe": "FFEB9C",
    "Devolución con pedido BC (comprobar abono)": "D9C3EA",
    "Pendiente en BC (sin procesar)": "FFF2CC",
    "Error BC (sin pedido BC)": "F8CBAD",
    "No es venta: cancelado / rechazado en Mirakl": "DDEBF7",
    "Revisar: pedido BC creado para línea cancelada/rechazada": "FFC000",
    "Solo en export Mirakl (falta en el log BC)": "F4B183",
    "Solo en log BC (no está en el export)": "FFC7CE",
}
EXPLICACION = {
    "Coincide: pedido BC creado": "Está en el export y BC creó el pedido de venta (Processed) con el mismo importe.",
    "Coincide con diferencia de importe": "Está en ambos y BC creó el pedido, pero el importe del log no coincide con el del export.",
    "Devolución con pedido BC (comprobar abono)": "Mirakl la da como reembolsada o con incidencia de devolución y BC había creado el pedido: comprobar abono/devolución en BC.",
    "Pendiente en BC (sin procesar)": "Está en el log del conector con estado Pending: BC todavía no ha creado nada para esta línea.",
    "Error BC (sin pedido BC)": "El conector falló al crear el pedido (ver Error BC): la venta falta en BC.",
    "No es venta: cancelado / rechazado en Mirakl": "La línea se canceló o se rechazó en Mirakl: no debe haber pedido en BC.",
    "Revisar: pedido BC creado para línea cancelada/rechazada": "BC creó pedido para una línea que Mirakl da como cancelada o rechazada.",
    "Solo en export Mirakl (falta en el log BC)": "Está en el export de Mirakl y no aparece en el log del conector.",
    "Solo en log BC (no está en el export)": "Está en el log del conector pero no en el export de Mirakl.",
}
COLOR_DEV = {"Sí": "C9A9E0", "Incidencia": "E4DFEC"}
COLOR_SIT = {
    "Pedido BC creado: comprobar abono/devolución en BC": "D9C3EA",
    "Error BC: sin pedido en BC": "F8CBAD",
    "Pendiente en BC: sin pedido ni abono todavía": "FFF2CC",
    "Sin pedido en BC": "F4B183",
}
MAP_ESTADO = {
    "RECEIVED": "Recibido", "REFUNDED": "Reembolsado", "SHIPPED": "Enviado", "INCIDENT_OPEN": "Incidencia abierta",
    "SHIPPING": "Esperando envío", "TO_COLLECT": "Para recoger", "REFUSED": "Rechazado", "CANCELED": "Cancelado",
    "WAITING_DEBIT_PAYMENT": "Débito en curso", "WAITING_ACCEPTANCE": "Esperando aceptación", "CLOSED": "Cerrado",
    "WAITING_REFUND_PAYMENT": "Esperando reembolso", "STAGING": "Preparación",
}

CAB_FILL = PatternFill("solid", fgColor="1F4E78")
CAB_FONT = Font(name=FUENTE, size=10, bold=True, color="FFFFFF")
NORMAL = Font(name=FUENTE, size=10)
BOLD = Font(name=FUENTE, size=10, bold=True)
TITULO = Font(name=FUENTE, size=14, bold=True, color="1F4E78")
NOTA = Font(name=FUENTE, size=9, italic=True, color="595959")
AZUL = Font(name=FUENTE, size=10, color="0000FF")
THIN = Side(style="thin", color="BFBFBF")
BORDE = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FMT_FECHA = "dd/mm/yyyy hh:mm"
FMT_DIA = "dd/mm/yyyy"
FMT_EUR = '#,##0.00 "€"'
FMT_DIF = '#,##0.00 "€";[Red]-#,##0.00 "€";"-"'
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
         "septiembre", "octubre", "noviembre", "diciembre"]


# ----------------------------------------------------------------------------- lectura
def _iso_madrid(s: pd.Series) -> pd.Series:
    """'2026-06-01T02:14:21Z' -> datetime naive en hora de Madrid."""
    return pd.to_datetime(s.replace("", None), errors="coerce", utc=True).dt.tz_convert("Europe/Madrid").dt.tz_localize(None)


def leer_log_bc(ruta: Path) -> pd.DataFrame:
    df = pd.read_excel(ruta, dtype=str, keep_default_na=False)
    df.columns = [str(c).strip() for c in df.columns]
    for c in ["Mirakl Line Id", "Order ID", "Processing Status"]:
        if c not in df.columns:
            sys.exit(f"El log BC no tiene la columna '{c}'")
    df = df[df["Mirakl Line Id"].str.strip() != ""].copy()
    df["Fecha creación"] = _iso_madrid(df["Created Date"])
    df["Fecha envío (log)"] = _iso_madrid(df["Shipped Date"])
    df["Fecha reembolso (log)"] = _iso_madrid(df["Refund Created Date"])
    df["Estado línea (ES)"] = df["Line State"].map(lambda v: MAP_ESTADO.get(v, v))
    df["Estado pedido (ES)"] = df["Order State"].map(lambda v: MAP_ESTADO.get(v, v))
    for c in ["Quantity", "Price", "Total Price", "Refund Amount", "Return Qty", "Return Qty Detected"]:
        df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0)
    if df["Mirakl Line Id"].duplicated().any():
        print("AVISO: Mirakl Line Id duplicados en el log:", df.loc[df["Mirakl Line Id"].duplicated(), "Mirakl Line Id"].tolist()[:10])
    return df.reset_index(drop=True)


def leer_export_mirakl(ruta: Path) -> pd.DataFrame:
    df = pd.read_excel(ruta, sheet_name="orders", dtype=str, keep_default_na=False)
    df.columns = [str(c).strip() for c in df.columns]
    for c in ["N.º de asiento de pedido", "Número de pedido", "Estado", "Importe"]:
        if c not in df.columns:
            sys.exit(f"El export de Mirakl no tiene la columna '{c}'")
    df = df[df["N.º de asiento de pedido"].str.strip() != ""].copy()
    df["Fecha creación"] = pd.to_datetime(df["Fecha de creación"], format="%d/%m/%Y %H:%M:%S", errors="coerce")
    for c in ["Fecha de aceptación", "Fecha límite de envío", "Fecha de envío", "Fecha de recepción", "Entrega"]:
        df[c] = pd.to_datetime(df[c].replace("", None), format="%d/%m/%Y %H:%M:%S", errors="coerce")
    for c in ["Fecha Emisión", "Fecha de cumplimentación", "Fecha de abono 1", "Fecha de abono 2", "Fecha de recogida 1", "Fecha de recogida 2"]:
        df[c] = pd.to_datetime(df[c].replace("", None), format="%Y%m%d", errors="coerce")
    for c in ["Cantidad", "Precio por unidad", "Importe", "Importe total del pedido con IVA (gastos de envío incluidos)",
              "Comisión (sin impuestos)", "Valor de la comisión (impuestos incluidos)", "Importe transferido a tienda (impuestos incluidos)",
              "Importe total cancelado (impuestos incluidos)", "Importe total reembolsado (impuestos incluidos)"]:
        df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0)
    if df["N.º de asiento de pedido"].duplicated().any():
        print("AVISO: asientos duplicados en el export:", df.loc[df["N.º de asiento de pedido"].duplicated(), "N.º de asiento de pedido"].tolist()[:10])
    return df.reset_index(drop=True)


# ----------------------------------------------------------------------------- escritura
def _cabecera(ws, fila, cabeceras, anchos=None):
    for j, h in enumerate(cabeceras, start=1):
        c = ws.cell(row=fila, column=j, value=h)
        c.fill, c.font, c.border = CAB_FILL, CAB_FONT, BORDE
        c.alignment = Alignment(vertical="center", wrap_text=True)
        if anchos and j - 1 < len(anchos) and anchos[j - 1]:
            ws.column_dimensions[get_column_letter(j)].width = anchos[j - 1]
    ws.row_dimensions[fila].height = 30


def _celda(ws, fila, col, valor, fmt=None, font=NORMAL, borde=True):
    if isinstance(valor, float) and pd.isna(valor):
        valor = None
    elif isinstance(valor, pd.Timestamp):
        valor = None if pd.isna(valor) else valor.to_pydatetime()
    elif valor is pd.NaT:
        valor = None
    c = ws.cell(row=fila, column=col, value=valor)
    c.font = font
    if borde:
        c.border = BORDE
    if fmt:
        c.number_format = fmt
    return c


def _pinta_filas(ws, col_clave, primera, ultima, ncols, colores):
    rango = f"A{primera}:{get_column_letter(ncols)}{ultima}"
    letra = get_column_letter(col_clave)
    for texto, rgb in colores.items():
        ws.conditional_formatting.add(
            rango, FormulaRule(formula=[f'${letra}{primera}="{texto}"'], fill=PatternFill("solid", fgColor=rgb, bgColor=rgb)))


def _escribe_df(ws, df, fila_cab, anchos, formatos):
    _cabecera(ws, fila_cab, list(df.columns), anchos)
    cols = list(df.columns)
    for i, fila in enumerate(df.itertuples(index=False), start=fila_cab + 1):
        for j, v in enumerate(fila, start=1):
            _celda(ws, i, j, v, formatos.get(cols[j - 1]))
    ws.freeze_panes = ws.cell(row=fila_cab + 1, column=2)
    ws.auto_filter.ref = f"A{fila_cab}:{get_column_letter(len(cols))}{max(fila_cab + 1, fila_cab + len(df))}"


def _tabla_cabecera(ws, r, textos):
    _cabecera(ws, r, textos)
    ws.cell(row=r, column=1).fill = PatternFill(fill_type=None)
    ws.cell(row=r, column=1).border = Border()


def _sin_avisos_numero_texto(ruta: Path):
    tag = f'<ignoredErrors><ignoredError sqref="A1:BZ{MAXF}" numberStoredAsText="1"/></ignoredErrors>'
    tmp = ruta.with_suffix(".tmp.xlsx")
    with zipfile.ZipFile(ruta) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            datos = zin.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet") and item.filename.endswith(".xml"):
                xml = datos.decode("utf-8")
                if "<ignoredErrors" not in xml and "</worksheet>" in xml and "<tableParts" not in xml and "<drawing" not in xml:
                    datos = xml.replace("</worksheet>", tag + "</worksheet>").encode("utf-8")
            zout.writestr(item, datos)
    shutil.move(tmp, ruta)


COLS_EX = [
    ("N.º de asiento de pedido", "N.º asiento"), ("Número de pedido", "Nº pedido"), ("Fecha creación", "Fecha creación"),
    ("Canal", "Canal"), ("Estado", "Estado"), ("Motivo", "Motivo"), ("Cantidad", "Cantidad"), ("SKU del producto", "SKU (EAN)"),
    ("SKU de la oferta", "SKU oferta"), ("Detalles", "Detalles"), ("Precio por unidad", "Precio unidad"), ("Importe", "Importe"),
    ("Importe total del pedido con IVA (gastos de envío incluidos)", "Total pedido con IVA"),
    ("Comisión (sin impuestos)", "Comisión sin imp."), ("Valor de la comisión (impuestos incluidos)", "Comisión con imp."),
    ("Importe transferido a tienda (impuestos incluidos)", "Transferido a tienda"),
    ("Asiento de pedido con cancelaciones", "Con cancelación"), ("Importe total cancelado (impuestos incluidos)", "Importe cancelado"),
    ("Asiento de pedido con reembolsos", "Con reembolso"), ("Importe total reembolsado (impuestos incluidos)", "Importe reembolsado"),
    ("Asiento de pedido con incidentes", "Con incidente"), ("Método de envío", "Método de envío"), ("Método de pago", "Método de pago"),
    ("Dirección de entrega: país", "País entrega"), ("Fecha de aceptación", "Fecha aceptación"), ("Fecha límite de envío", "Fecha límite envío"),
    ("Fecha de envío", "Fecha envío"), ("Fecha de recepción", "Fecha recepción"), ("Entrega", "Entrega"),
    ("Empresa transportista", "Transportista"), ("Número de seguimiento", "Nº seguimiento"), ("Talón de venta", "Talón de venta"),
    ("Fecha Emisión", "Fecha emisión"), ("Código de cesta", "Código de cesta"), ("Fecha de cumplimentación", "Fecha cumplimentación"),
    ("Talón de cumplimentación", "Talón cumplimentación"), ("Fecha de abono 1", "Fecha abono 1"), ("Talón de abono 1", "Talón abono 1"),
    ("Fecha de abono 2", "Fecha abono 2"), ("Talón de abono 2", "Talón abono 2"), ("Fecha de recogida 1", "Fecha recogida 1"),
    ("Talón de recogida 1", "Talón recogida 1"), ("Fecha de recogida 2", "Fecha recogida 2"), ("Talón de recogida 2", "Talón recogida 2"),
    ("stl-ovdev", "stl-ovdev"),
]
COLS_LOG = [
    "Mirakl Line Id", "Order ID", "Commercial ID", "Fecha creación", "Order State", "Estado pedido (ES)", "Line State", "Estado línea (ES)",
    "Channel", "SKU", "Title", "Quantity", "Price", "Total Price", "Fecha envío (log)", "Country", "City", "Zip Code",
    "Store Pickup Required", "ECI Label Status", "Processing Status", "BC Sales Order No.", "Error Message", "Item Found",
    "Return detected", "Return Qty Detected", "Return Status", "Return Source", "Return Qty", "Return Processing Status",
    "BC Return Order No.", "Reason Code", "Refund Id", "Order Refund Id", "Refund Amount", "Fecha reembolso (log)", "Return Request Id",
]


def construir(log, ex, periodo, salida, ruta_log, ruta_ex):
    anyo, mes = (int(x) for x in periodo.split("-"))
    ini = dt.datetime(anyo, mes, 1)
    fin = dt.datetime(anyo, mes, calendar.monthrange(anyo, mes)[1], 23, 59, 59)
    nombre_mes = MESES[mes - 1]

    la = log[(log["Fecha creación"] >= ini) & (log["Fecha creación"] <= fin)].copy()
    ea = ex[(ex["Fecha creación"] >= ini) & (ex["Fecha creación"] <= fin)].copy()
    ea = ea.sort_values("Fecha creación", ascending=False).reset_index(drop=True)
    la = la.sort_values("Fecha creación", ascending=False).reset_index(drop=True)
    pivot_log = pd.crosstab(log["Fecha creación"].dt.strftime("%Y-%m"), log["Processing Status"])

    # --- unión de líneas ------------------------------------------------------------
    idx_ex = ea.set_index("N.º de asiento de pedido")
    idx_log = la.set_index("Mirakl Line Id")
    union = []
    for k, r in idx_ex.iterrows():
        union.append((k, r["Número de pedido"], r["Fecha creación"], r["Canal"], r["SKU del producto"], str(r["Detalles"])[:70], r["Cantidad"]))
    for k, r in idx_log.iterrows():
        if k not in idx_ex.index:
            union.append((k, r["Order ID"], r["Fecha creación"], r["Channel"], r["SKU"], str(r["Title"])[:70], r["Quantity"]))
    union.sort(key=lambda t: (pd.notna(t[2]), t[2] if pd.notna(t[2]) else dt.datetime.min), reverse=True)
    pedidos = []
    vistos = set()
    for k, ped, f, canal, *_ in union:
        if ped not in vistos:
            vistos.add(ped)
            pedidos.append((ped, f, canal))

    wb = openpyxl.Workbook()
    ws_res = wb.active
    ws_res.title = "Resumen"
    ws_cl = wb.create_sheet("Cruce líneas")
    ws_cp = wb.create_sheet("Cruce pedidos")
    ws_dev = wb.create_sheet("Devoluciones")
    ws_ex = wb.create_sheet("Export Mirakl")
    ws_log = wb.create_sheet("Log BC")

    # --- hoja Export Mirakl ---------------------------------------------------------
    pex = ea[[a for a, _ in COLS_EX]].copy()
    pex.columns = [b for _, b in COLS_EX]
    ws_ex["A1"] = f"Export de pedidos de ECI / Mirakl ({ruta_ex.name}, hoja 'orders'), líneas con fecha de creación en {nombre_mes} de {anyo}"
    ws_ex["A1"].font = TITULO
    ws_ex["A2"] = "Solo columnas de negocio: sin nombres, direcciones ni teléfonos del cliente. Fechas en hora de Madrid, como las exporta Mirakl."
    ws_ex["A2"].font = NOTA
    FE = 4
    fmt_ex = {"Fecha creación": FMT_FECHA, "Precio unidad": FMT_EUR, "Importe": FMT_EUR, "Total pedido con IVA": FMT_EUR,
              "Comisión sin imp.": FMT_EUR, "Comisión con imp.": FMT_EUR, "Transferido a tienda": FMT_EUR, "Importe cancelado": FMT_EUR,
              "Importe reembolsado": FMT_EUR, "Fecha aceptación": FMT_FECHA, "Fecha límite envío": FMT_FECHA, "Fecha envío": FMT_FECHA,
              "Fecha recepción": FMT_FECHA, "Entrega": FMT_FECHA, "Fecha emisión": FMT_DIA, "Fecha cumplimentación": FMT_DIA,
              "Fecha abono 1": FMT_DIA, "Fecha abono 2": FMT_DIA, "Fecha recogida 1": FMT_DIA, "Fecha recogida 2": FMT_DIA}
    _escribe_df(ws_ex, pex, FE, [34, 32, 16, 14, 16, 14, 8, 15, 24, 40, 11, 11, 13, 11, 11, 13, 10, 11, 10, 12, 10, 22, 9, 8,
                                 16, 16, 16, 16, 16, 16, 24, 12, 12, 14, 12, 12, 12, 12, 12, 12, 12, 12, 12, 12, 28], fmt_ex)
    EX = {c: get_column_letter(i + 1) for i, c in enumerate(pex.columns)}

    def rex(col):
        return f"'Export Mirakl'!${EX[col]}${FE + 1}:${EX[col]}${MAXF}"

    # --- hoja Log BC ----------------------------------------------------------------
    plog = la[COLS_LOG].copy()
    ws_log["A1"] = f"Log del conector Mirakl → Business Central ({ruta_log.name}), líneas con fecha de creación en {nombre_mes} de {anyo}"
    ws_log["A1"].font = TITULO
    ws_log["A2"] = ("Sin nombre del cliente. 'Fecha creación' es Created Date pasada de UTC a hora de Madrid. "
                    "Los estados son los del momento en que se extrajo el log, no los actuales.")
    ws_log["A2"].font = NOTA
    FL = 4
    _escribe_df(ws_log, plog, FL, [36, 34, 32, 16, 12, 14, 14, 16, 14, 15, 40, 8, 10, 11, 16, 7, 18, 8, 8, 10, 12, 14, 60, 7, 7, 7, 14, 10, 7, 12, 12, 10, 10, 34, 11, 16, 30],
                {"Fecha creación": FMT_FECHA, "Price": FMT_EUR, "Total Price": FMT_EUR, "Fecha envío (log)": FMT_FECHA,
                 "Refund Amount": FMT_EUR, "Fecha reembolso (log)": FMT_FECHA})
    LG = {c: get_column_letter(i + 1) for i, c in enumerate(plog.columns)}

    def rlog(col):
        return f"'Log BC'!${LG[col]}${FL + 1}:${LG[col]}${MAXF}"

    # --- hoja Cruce líneas ----------------------------------------------------------
    ws_cl["A1"] = f"Cruce de líneas de pedido ECI (Mirakl) vs Business Central — {nombre_mes} {anyo}"
    ws_cl["A1"].font = TITULO
    ws_cl["A2"] = ("Una fila por línea de pedido (N.º de asiento = Mirakl Line Id), unión del export de Mirakl y del log del conector. "
                   "Las columnas H en adelante son fórmulas sobre 'Export Mirakl' y 'Log BC'; el color de la fila sigue Resultado y la columna Devolución va en lila.")
    ws_cl["A2"].font = NOTA
    cab = ["N.º asiento (clave)", "Nº pedido Mirakl", "Fecha creación", "Canal", "SKU (EAN)", "Descripción", "Cantidad",
           "En export", "Estado Mirakl (export)", "Motivo", "Importe Mirakl", "Reembolso Mirakl",
           "En log BC", "Estado línea (log BC)", "Importe log BC", "Dif. importe (log − export)",
           "Estado proceso BC", "Nº pedido BC", "Error BC", "Estado devolución (log BC)", "Proceso devolución BC", "Nº devolución BC",
           "Estado cambiado desde el log", "Devolución", "Resultado"]
    FCL = 4
    _cabecera(ws_cl, FCL, cab, [36, 32, 16, 14, 15, 44, 8, 8, 16, 12, 12, 12, 8, 16, 12, 12, 12, 12, 44, 14, 12, 12, 11, 11, 44])
    L = {c: get_column_letter(i + 1) for i, c in enumerate(cab)}
    p1 = FCL + 1
    for i, (k, ped, f, canal, sku, desc, qty) in enumerate(union, start=p1):
        _celda(ws_cl, i, 1, k)
        _celda(ws_cl, i, 2, ped)
        _celda(ws_cl, i, 3, f, FMT_FECHA)
        _celda(ws_cl, i, 4, canal)
        _celda(ws_cl, i, 5, sku)
        _celda(ws_cl, i, 6, desc)
        _celda(ws_cl, i, 7, float(qty) if pd.notna(qty) else None, "0")
        A = f"$A{i}"
        e, h = f'{L["En export"]}{i}', f'{L["En log BC"]}{i}'
        mex = lambda col: f'INDEX({rex(col)},MATCH({A},{rex("N.º asiento")},0))'
        mlog = lambda col: f'INDEX({rlog(col)},MATCH({A},{rlog("Mirakl Line Id")},0))'
        est, dev_ex = f'{L["Estado Mirakl (export)"]}{i}', f'{L["Reembolso Mirakl"]}{i}'
        q, r_, t = f'{L["Estado proceso BC"]}{i}', f'{L["Nº pedido BC"]}{i}', f'{L["Estado devolución (log BC)"]}{i}'
        x, pdif = f'{L["Devolución"]}{i}', f'{L["Dif. importe (log − export)"]}{i}'
        formulas = {
            "En export": f'=IF(COUNTIF({rex("N.º asiento")},{A})>0,"Sí","No")',
            "Estado Mirakl (export)": f'=IF({e}="Sí",{mex("Estado")}&"","")',
            "Motivo": f'=IF({e}="Sí",{mex("Motivo")}&"","")',
            "Importe Mirakl": f'=IF({e}="Sí",SUMIF({rex("N.º asiento")},{A},{rex("Importe")}),"")',
            "Reembolso Mirakl": f'=IF({e}="Sí",SUMIF({rex("N.º asiento")},{A},{rex("Importe reembolsado")}),"")',
            "En log BC": f'=IF(COUNTIF({rlog("Mirakl Line Id")},{A})>0,"Sí","No")',
            "Estado línea (log BC)": f'=IF({h}="Sí",{mlog("Estado línea (ES)")}&"","")',
            "Importe log BC": f'=IF({h}="Sí",SUMIF({rlog("Mirakl Line Id")},{A},{rlog("Total Price")}),"")',
            "Dif. importe (log − export)": f'=IF(AND({e}="Sí",{h}="Sí"),ROUND({L["Importe log BC"]}{i}-{L["Importe Mirakl"]}{i},2),"")',
            "Estado proceso BC": f'=IF({h}="Sí",{mlog("Processing Status")}&"","")',
            "Nº pedido BC": f'=IF({h}="Sí",{mlog("BC Sales Order No.")}&"","")',
            "Error BC": f'=IF({h}="Sí",{mlog("Error Message")}&"","")',
            "Estado devolución (log BC)": f'=IF({h}="Sí",{mlog("Return Status")}&"","")',
            "Proceso devolución BC": f'=IF({h}="Sí",{mlog("Return Processing Status")}&"","")',
            "Nº devolución BC": f'=IF({h}="Sí",{mlog("BC Return Order No.")}&"","")',
            "Estado cambiado desde el log": f'=IF(AND({e}="Sí",{h}="Sí"),IF({est}<>{L["Estado línea (log BC)"]}{i},"Sí","No"),"")',
            "Devolución": (f'=IF(OR({est}="Reembolsado",AND({dev_ex}<>"",{dev_ex}>0)),"Sí",'
                           f'IF({est}="Incidencia abierta","Incidencia",'
                           f'IF(AND({e}="No",{t}="REFUNDED"),"Sí",IF(AND({e}="No",{t}="INCIDENT_OPEN"),"Incidencia","No"))))'),
            "Resultado": (f'=IF(AND({h}="No",{e}="Sí"),"Solo en export Mirakl (falta en el log BC)",'
                          f'IF({e}="No","Solo en log BC (no está en el export)",'
                          f'IF(AND(OR({est}="Cancelado",{est}="Rechazado"),{r_}<>""),"Revisar: pedido BC creado para línea cancelada/rechazada",'
                          f'IF(OR({est}="Cancelado",{est}="Rechazado"),"No es venta: cancelado / rechazado en Mirakl",'
                          f'IF({q}="Error","Error BC (sin pedido BC)",'
                          f'IF({q}="Pending","Pendiente en BC (sin procesar)",'
                          f'IF({x}<>"No","Devolución con pedido BC (comprobar abono)",'
                          f'IF(AND({pdif}<>"",{pdif}<>0),"Coincide con diferencia de importe","Coincide: pedido BC creado"))))))))'),
        }
        fmts = {"Importe Mirakl": FMT_EUR, "Reembolso Mirakl": FMT_EUR, "Importe log BC": FMT_EUR, "Dif. importe (log − export)": FMT_DIF}
        for col, fx in formulas.items():
            _celda(ws_cl, i, cab.index(col) + 1, fx, fmts.get(col))
    pN = p1 + len(union) - 1
    ws_cl.freeze_panes = f"B{p1}"
    ws_cl.auto_filter.ref = f"A{FCL}:{get_column_letter(len(cab))}{pN}"
    cD = L["Devolución"]
    for texto, rgb in COLOR_DEV.items():
        ws_cl.conditional_formatting.add(
            f"{cD}{p1}:{cD}{pN}", FormulaRule(formula=[f'${cD}{p1}="{texto}"'], fill=PatternFill("solid", fgColor=rgb, bgColor=rgb),
                                              font=Font(name=FUENTE, size=10, bold=True)))
    _pinta_filas(ws_cl, cab.index("Resultado") + 1, p1, pN, len(cab), COLOR)

    def rcl(col):
        return f"'Cruce líneas'!${L[col]}${p1}:${L[col]}${pN}"

    # --- hoja Cruce pedidos ---------------------------------------------------------
    ws_cp["A1"] = f"Cruce por pedido ECI (Mirakl) vs Business Central — {nombre_mes} {anyo}"
    ws_cp["A1"].font = TITULO
    ws_cp["A2"] = "Una fila por pedido de Mirakl, agregando las líneas de la hoja 'Cruce líneas' con fórmulas. El color de la fila sigue 'Resultado pedido'."
    ws_cp["A2"].font = NOTA
    cab_p = ["Nº pedido Mirakl", "Fecha creación", "Canal", "Líneas", "Importe Mirakl", "Importe log BC", "Dif. importe",
             "Estado pedido (log BC)", "Estado proceso BC", "Nº pedido BC", "Líneas devueltas", "Líneas con incidencia",
             "Importe reembolsado", "Líneas canceladas / rechazadas", "Resultado pedido"]
    FCP = 4
    _cabecera(ws_cp, FCP, cab_p, [32, 16, 14, 7, 13, 13, 12, 16, 12, 12, 9, 9, 13, 11, 40])
    P = {c: get_column_letter(i + 1) for i, c in enumerate(cab_p)}
    q1 = FCP + 1
    for i, (ped, f, canal) in enumerate(pedidos, start=q1):
        _celda(ws_cp, i, 1, ped)
        _celda(ws_cp, i, 2, f, FMT_FECHA)
        _celda(ws_cp, i, 3, canal)
        A = f"$A{i}"
        cnt = lambda col, val: f'COUNTIFS({rcl("Nº pedido Mirakl")},{A},{rcl(col)},"{val}")'
        d, ii, n = f'{P["Líneas"]}{i}', f'{P["Estado proceso BC"]}{i}', f'{P["Líneas canceladas / rechazadas"]}{i}'
        k_, l_ = f'{P["Líneas devueltas"]}{i}', f'{P["Líneas con incidencia"]}{i}'
        fx = {
            "Líneas": f'=COUNTIF({rcl("Nº pedido Mirakl")},{A})',
            "Importe Mirakl": f'=SUMIF({rcl("Nº pedido Mirakl")},{A},{rcl("Importe Mirakl")})',
            "Importe log BC": f'=SUMIF({rcl("Nº pedido Mirakl")},{A},{rcl("Importe log BC")})',
            "Dif. importe": f'=ROUND({P["Importe log BC"]}{i}-{P["Importe Mirakl"]}{i},2)',
            "Estado pedido (log BC)": f'=IFERROR(INDEX({rlog("Estado pedido (ES)")},MATCH({A},{rlog("Order ID")},0))&"","")',
            "Estado proceso BC": (f'=IF({cnt("Estado proceso BC", "Error")}>0,"Error",IF({cnt("Estado proceso BC", "Pending")}>0,"Pending",'
                                  f'IF({cnt("Estado proceso BC", "Processed")}>0,"Processed","")))'),
            "Nº pedido BC": f'=IFERROR(INDEX({rlog("BC Sales Order No.")},MATCH({A},{rlog("Order ID")},0))&"","")',
            "Líneas devueltas": f'={cnt("Devolución", "Sí")}',
            "Líneas con incidencia": f'={cnt("Devolución", "Incidencia")}',
            "Importe reembolsado": f'=SUMIF({rcl("Nº pedido Mirakl")},{A},{rcl("Reembolso Mirakl")})',
            "Líneas canceladas / rechazadas": f'={cnt("Estado Mirakl (export)", "Cancelado")}+{cnt("Estado Mirakl (export)", "Rechazado")}',
            "Resultado pedido": (f'=IF({cnt("Resultado", "Solo en export*")}>0,"Solo en export Mirakl (falta en el log BC)",'
                                 f'IF({cnt("Resultado", "Solo en log*")}>0,"Solo en log BC (no está en el export)",'
                                 f'IF({cnt("Resultado", "Revisar*")}>0,"Revisar: pedido BC creado para línea cancelada/rechazada",'
                                 f'IF({n}={d},"No es venta: cancelado / rechazado en Mirakl",'
                                 f'IF({ii}="Error","Error BC (sin pedido BC)",IF({ii}="Pending","Pendiente en BC (sin procesar)",'
                                 f'IF({k_}+{l_}>0,"Devolución con pedido BC (comprobar abono)",'
                                 f'IF({P["Dif. importe"]}{i}<>0,"Coincide con diferencia de importe","Coincide: pedido BC creado"))))))))'),
        }
        fmts = {"Importe Mirakl": FMT_EUR, "Importe log BC": FMT_EUR, "Dif. importe": FMT_DIF, "Importe reembolsado": FMT_EUR}
        for col, f_ in fx.items():
            _celda(ws_cp, i, cab_p.index(col) + 1, f_, fmts.get(col))
    qN = q1 + len(pedidos) - 1
    ws_cp.freeze_panes = f"B{q1}"
    ws_cp.auto_filter.ref = f"A{FCP}:{get_column_letter(len(cab_p))}{qN}"
    _pinta_filas(ws_cp, cab_p.index("Resultado pedido") + 1, q1, qN, len(cab_p), COLOR)

    def rcp(col):
        return f"'Cruce pedidos'!${P[col]}${q1}:${P[col]}${qN}"

    # --- hoja Devoluciones ----------------------------------------------------------
    def es_devolucion(k):
        if k in idx_ex.index:
            r = idx_ex.loc[k]
            return r["Estado"] in ("Reembolsado", "Incidencia abierta") or r["Importe total reembolsado (impuestos incluidos)"] > 0
        r = idx_log.loc[k]
        return r["Return Status"] in ("REFUNDED", "INCIDENT_OPEN")

    devs = [u for u in union if es_devolucion(u[0])]
    ws_dev["A1"] = f"Devoluciones y reembolsos ECI — {nombre_mes} {anyo}"
    ws_dev["A1"].font = TITULO
    ws_dev["A2"] = ("Líneas con Devolución = Sí (reembolsadas en Mirakl) o Incidencia (devolución en curso). Datos del reembolso del export "
                    "(importe, talón y fecha de abono y de recogida de ECI) y del log BC (Return Status, motivo, Refund Id). El color sigue 'Situación en BC'.")
    ws_dev["A2"].font = NOTA
    cab_d = ["N.º asiento (clave)", "Nº pedido Mirakl", "Fecha creación", "Canal", "SKU (EAN)", "Descripción", "Cantidad",
             "Devolución", "Estado Mirakl (export)", "Motivo", "Importe línea", "Importe reembolsado", "Reembolso parcial",
             "Fecha abono 1", "Talón abono 1", "Fecha recogida 1", "Talón recogida 1",
             "Estado devolución (log BC)", "Motivo (log BC)", "Refund Id (log)", "Importe reembolso (log)", "Fecha reembolso (log)",
             "Estado proceso BC", "Nº pedido BC", "Proceso devolución BC", "Nº devolución BC", "Situación en BC"]
    FCD = 4
    _cabecera(ws_dev, FCD, cab_d, [36, 32, 16, 14, 15, 40, 8, 10, 16, 12, 12, 12, 9, 12, 12, 12, 12, 14, 11, 11, 12, 16, 12, 12, 12, 12, 44])
    D = {c: get_column_letter(i + 1) for i, c in enumerate(cab_d)}
    d1 = FCD + 1
    for i, (k, ped, f, canal, sku, desc, qty) in enumerate(devs, start=d1):
        _celda(ws_dev, i, 1, k)
        _celda(ws_dev, i, 2, ped)
        _celda(ws_dev, i, 3, f, FMT_FECHA)
        _celda(ws_dev, i, 4, canal)
        _celda(ws_dev, i, 5, sku)
        _celda(ws_dev, i, 6, desc)
        _celda(ws_dev, i, 7, float(qty) if pd.notna(qty) else None, "0")
        A = f"$A{i}"
        mcl = lambda col: f'INDEX({rcl(col)},MATCH({A},{rcl("N.º asiento (clave)")},0))'
        mex = lambda col: f'IFERROR(INDEX({rex(col)},MATCH({A},{rex("N.º asiento")},0)),"")'
        mlog = lambda col: f'IFERROR(INDEX({rlog(col)},MATCH({A},{rlog("Mirakl Line Id")},0)),"")'
        w_, v_ = f'{D["Nº pedido BC"]}{i}', f'{D["Estado proceso BC"]}{i}'
        fx = {
            "Devolución": f'={mcl("Devolución")}&""',
            "Estado Mirakl (export)": f'={mcl("Estado Mirakl (export)")}&""',
            "Motivo": f'={mcl("Motivo")}&""',
            "Importe línea": f'=IF({mcl("Importe Mirakl")}="",{mcl("Importe log BC")},{mcl("Importe Mirakl")})',
            "Importe reembolsado": f'={mcl("Reembolso Mirakl")}',
            "Reembolso parcial": (f'=IF(AND({D["Importe reembolsado"]}{i}<>"",{D["Importe reembolsado"]}{i}>0,'
                                  f'{D["Importe reembolsado"]}{i}<{D["Importe línea"]}{i}-0.005),"Sí","No")'),
            "Fecha abono 1": f'=IF({mex("Fecha abono 1")}="","",{mex("Fecha abono 1")})',
            "Talón abono 1": f'={mex("Talón abono 1")}&""',
            "Fecha recogida 1": f'=IF({mex("Fecha recogida 1")}="","",{mex("Fecha recogida 1")})',
            "Talón recogida 1": f'={mex("Talón recogida 1")}&""',
            "Estado devolución (log BC)": f'={mcl("Estado devolución (log BC)")}&""',
            "Motivo (log BC)": f'={mlog("Reason Code")}&""',
            "Refund Id (log)": f'={mlog("Refund Id")}&""',
            "Importe reembolso (log)": f'=IF({mlog("Refund Amount")}="","",{mlog("Refund Amount")})',
            "Fecha reembolso (log)": f'=IF({mlog("Fecha reembolso (log)")}="","",{mlog("Fecha reembolso (log)")})',
            "Estado proceso BC": f'={mcl("Estado proceso BC")}&""',
            "Nº pedido BC": f'={mcl("Nº pedido BC")}&""',
            "Proceso devolución BC": f'={mcl("Proceso devolución BC")}&""',
            "Nº devolución BC": f'={mcl("Nº devolución BC")}&""',
            "Situación en BC": (f'=IF({w_}<>"","Pedido BC creado: comprobar abono/devolución en BC",'
                                f'IF({v_}="Error","Error BC: sin pedido en BC",'
                                f'IF({v_}="Pending","Pendiente en BC: sin pedido ni abono todavía","Sin pedido en BC")))'),
        }
        fmts = {"Importe línea": FMT_EUR, "Importe reembolsado": FMT_EUR, "Fecha abono 1": FMT_DIA, "Fecha recogida 1": FMT_DIA,
                "Importe reembolso (log)": FMT_EUR, "Fecha reembolso (log)": FMT_FECHA}
        for col, f_ in fx.items():
            _celda(ws_dev, i, cab_d.index(col) + 1, f_, fmts.get(col))
    dN = d1 + max(len(devs), 1) - 1
    ws_dev.freeze_panes = f"B{d1}"
    ws_dev.auto_filter.ref = f"A{FCD}:{get_column_letter(len(cab_d))}{dN}"
    cV = D["Devolución"]
    for texto, rgb in COLOR_DEV.items():
        ws_dev.conditional_formatting.add(
            f"{cV}{d1}:{cV}{dN}", FormulaRule(formula=[f'${cV}{d1}="{texto}"'], fill=PatternFill("solid", fgColor=rgb, bgColor=rgb),
                                              font=Font(name=FUENTE, size=10, bold=True)))
    _pinta_filas(ws_dev, cab_d.index("Situación en BC") + 1, d1, dN, len(cab_d), COLOR_SIT)

    def rdv(col):
        return f"'Devoluciones'!${D[col]}${d1}:${D[col]}${dN}"

    # --- hoja Resumen ---------------------------------------------------------------
    ws = ws_res
    for col, w in zip("ABCDEFG", [3, 64, 16, 18, 18, 18, 90]):
        ws.column_dimensions[col].width = w
    ws["B1"] = f"Cuadre ECI (Mirakl) vs Business Central — {nombre_mes} {anyo}"
    ws["B1"].font = TITULO
    ws["B2"] = (f"Generado el {dt.date.today():%d/%m/%Y} a partir de '{ruta_log.name}' (log del conector BC) y "
                f"'{ruta_ex.name}' (export de pedidos de Mirakl). Celdas en azul: parámetros.")
    ws["B2"].font = NOTA
    ws["B4"] = "Parámetros y cobertura"
    ws["B4"].font = BOLD
    filas_param = [
        ("Periodo: inicio (fecha de creación, hora de Madrid)", ini, FMT_DIA, AZUL, "Las hojas de datos y el cruce ya vienen filtradas por este periodo."),
        ("Periodo: fin", fin, FMT_FECHA, AZUL, ""),
        ("Log BC: líneas en total / en el periodo", f"{len(log)} / {len(la)}", None, NORMAL, f"Log completo del {log['Fecha creación'].min():%d/%m/%Y} al {log['Fecha creación'].max():%d/%m/%Y %H:%M}."),
        ("Log BC: última línea creada (≈ fecha de extracción del log)", log["Fecha creación"].max(), FMT_FECHA, NORMAL,
         "Los estados del log (línea, devolución, proceso) son los de ese momento."),
        ("Export Mirakl: líneas en total / en el periodo", f"{len(ex)} / {len(ea)}", None, NORMAL, f"Export completo del {ex['Fecha creación'].min():%d/%m/%Y} al {ex['Fecha creación'].max():%d/%m/%Y %H:%M}."),
        ("Export Mirakl: última línea creada (≈ fecha de extracción)", ex["Fecha creación"].max(), FMT_FECHA, NORMAL,
         "Los estados del export son los actuales a esa fecha: por eso pueden diferir de los del log."),
    ]
    r = 5
    for etiqueta, valor, fmt, fnt, nota in filas_param:
        _celda(ws, r, 2, etiqueta)
        _celda(ws, r, 3, valor, fmt, fnt)
        _celda(ws, r, 7, nota, font=NOTA, borde=False)
        r += 1
    r += 1
    ws.cell(row=r, column=2, value="Líneas y pedidos").font = BOLD
    r += 1
    _tabla_cabecera(ws, r, ["", "Concepto", "Líneas", "Pedidos", "Importe Mirakl", "Importe log BC", "Cómo se calcula"])
    r += 1
    kpis = [
        ("Líneas del export Mirakl en el periodo", f'=COUNTIF({rcl("En export")},"Sí")', f'=COUNTA({rcp("Nº pedido Mirakl")})',
         f'=SUMIF({rcl("En export")},"Sí",{rcl("Importe Mirakl")})', None, "Líneas del export con fecha de creación dentro del periodo. Pedidos = filas de 'Cruce pedidos'."),
        ("Líneas del log BC en el periodo", f'=COUNTIF({rcl("En log BC")},"Sí")', None, None,
         f'=SUMIF({rcl("En log BC")},"Sí",{rcl("Importe log BC")})', "Líneas del log del conector con fecha de creación (hora de Madrid) dentro del periodo."),
        ("   en ambos (export y log)", f'=COUNTIFS({rcl("En export")},"Sí",{rcl("En log BC")},"Sí")', None,
         f'=SUMIFS({rcl("Importe Mirakl")},{rcl("En export")},"Sí",{rcl("En log BC")},"Sí")',
         f'=SUMIFS({rcl("Importe log BC")},{rcl("En export")},"Sí",{rcl("En log BC")},"Sí")', ""),
        ("   diferencia de importe en las líneas en ambos", None, None, None,
         f'=SUMIFS({rcl("Importe log BC")},{rcl("En export")},"Sí",{rcl("En log BC")},"Sí")-SUMIFS({rcl("Importe Mirakl")},{rcl("En export")},"Sí",{rcl("En log BC")},"Sí")',
         "Importe log BC − importe Mirakl. Debe ser 0."),
        ("   con el estado cambiado entre el log y el export", f'=COUNTIF({rcl("Estado cambiado desde el log")},"Sí")', None, None, None,
         "Líneas cuyo estado en Mirakl ha cambiado desde que se extrajo el log (p. ej. Recibido → Reembolsado)."),
    ]
    for etiqueta, n, np_, imp_ex, imp_log, nota in kpis:
        _celda(ws, r, 2, etiqueta)
        _celda(ws, r, 3, n, "#,##0")
        _celda(ws, r, 4, np_, "#,##0")
        _celda(ws, r, 5, imp_ex, FMT_EUR)
        _celda(ws, r, 6, imp_log, FMT_EUR)
        _celda(ws, r, 7, nota, font=NOTA, borde=False)
        r += 1
    r += 1
    ws.cell(row=r, column=2, value="Qué ha hecho Business Central con las líneas del periodo").font = BOLD
    r += 1
    _tabla_cabecera(ws, r, ["", "Estado proceso BC", "Líneas", "Pedidos", "Importe Mirakl", "", "Cómo se calcula"])
    r += 1
    for est, nota in [("Processed", "Pedido de venta creado en BC (hay nº de pedido BC)."),
                      ("Error", "El conector intentó crearlo y falló: ver la columna Error BC."),
                      ("Pending", "El conector aún no lo ha procesado.")]:
        _celda(ws, r, 2, f"   {est}")
        _celda(ws, r, 3, f'=COUNTIF({rcl("Estado proceso BC")},"{est}")', "#,##0")
        _celda(ws, r, 4, f'=COUNTIF({rcp("Estado proceso BC")},"{est}")', "#,##0")
        _celda(ws, r, 5, f'=SUMIF({rcl("Estado proceso BC")},"{est}",{rcl("Importe Mirakl")})', FMT_EUR)
        _celda(ws, r, 6, None)
        _celda(ws, r, 7, nota, font=NOTA, borde=False)
        r += 1
    _celda(ws, r, 2, "   pedidos de venta BC distintos creados")
    _celda(ws, r, 3, None)
    _celda(ws, r, 4, f'=COUNTIF({rcp("Nº pedido BC")},"?*")', "#,##0")
    _celda(ws, r, 5, None)
    _celda(ws, r, 6, None)
    _celda(ws, r, 7, "Pedidos de 'Cruce pedidos' con nº de pedido BC.", font=NOTA, borde=False)
    r += 2

    ws.cell(row=r, column=2, value="Resultado del cruce por línea (color de fila en 'Cruce líneas' y 'Cruce pedidos')").font = BOLD
    r += 1
    _tabla_cabecera(ws, r, ["", "Resultado", "Líneas", "Pedidos", "Importe Mirakl", "Importe log BC", "Significado"])
    r += 1
    r_ini = r
    for texto, rgb in COLOR.items():
        c = _celda(ws, r, 2, texto)
        c.fill = PatternFill("solid", fgColor=rgb)
        _celda(ws, r, 3, f'=COUNTIF({rcl("Resultado")},"{texto}")', "#,##0")
        _celda(ws, r, 4, f'=COUNTIF({rcp("Resultado pedido")},"{texto}")', "#,##0")
        _celda(ws, r, 5, f'=SUMIF({rcl("Resultado")},"{texto}",{rcl("Importe Mirakl")})', FMT_EUR)
        _celda(ws, r, 6, f'=SUMIF({rcl("Resultado")},"{texto}",{rcl("Importe log BC")})', FMT_EUR)
        _celda(ws, r, 7, EXPLICACION[texto], font=NOTA, borde=False)
        r += 1
    _celda(ws, r, 2, "Total", font=BOLD)
    _celda(ws, r, 3, f"=SUM(C{r_ini}:C{r - 1})", "#,##0", BOLD)
    _celda(ws, r, 4, f"=SUM(D{r_ini}:D{r - 1})", "#,##0", BOLD)
    _celda(ws, r, 5, f"=SUM(E{r_ini}:E{r - 1})", FMT_EUR, BOLD)
    _celda(ws, r, 6, f"=SUM(F{r_ini}:F{r - 1})", FMT_EUR, BOLD)
    r += 2

    ws.cell(row=r, column=2, value="Devoluciones y reembolsos (hoja 'Devoluciones')").font = BOLD
    r += 1
    _tabla_cabecera(ws, r, ["", "Concepto", "Líneas", "Pedidos", "Importe línea", "Importe reembolsado", "Cómo se calcula"])
    r += 1
    devk = [
        ("Líneas reembolsadas en Mirakl (Devolución = Sí)", f'=COUNTIF({rcl("Devolución")},"Sí")', f'=COUNTIF({rcp("Líneas devueltas")},">0")',
         f'=SUMIF({rcl("Devolución")},"Sí",{rcl("Importe Mirakl")})', f'=SUMIF({rcl("Devolución")},"Sí",{rcl("Reembolso Mirakl")})',
         "Estado Mirakl = Reembolsado, o importe reembolsado > 0."),
        ("Líneas con incidencia abierta (devolución en curso)", f'=COUNTIF({rcl("Devolución")},"Incidencia")', f'=COUNTIF({rcp("Líneas con incidencia")},">0")',
         f'=SUMIF({rcl("Devolución")},"Incidencia",{rcl("Importe Mirakl")})', f'=SUMIF({rcl("Devolución")},"Incidencia",{rcl("Reembolso Mirakl")})',
         "Estado Mirakl = Incidencia abierta (motivo Devolución). Aún no reembolsadas, salvo parciales."),
        ("   reembolsos parciales (menos que el importe de la línea)", f'=COUNTIF({rdv("Reembolso parcial")},"Sí")', None, None, None,
         "Líneas de 2 unidades con 1 devuelta, normalmente."),
        ("Líneas canceladas o rechazadas en Mirakl (no son venta)", f'=COUNTIF({rcl("Estado Mirakl (export)")},"Cancelado")+COUNTIF({rcl("Estado Mirakl (export)")},"Rechazado")',
         f'=COUNTIF({rcp("Líneas canceladas / rechazadas")},">0")',
         f'=SUMIF({rcl("Estado Mirakl (export)")},"Cancelado",{rcl("Importe Mirakl")})+SUMIF({rcl("Estado Mirakl (export)")},"Rechazado",{rcl("Importe Mirakl")})', None,
         "Rechazado = rechazadas por la tienda (Toni Pons); Cancelado = canceladas por el cliente / no debitadas."),
        ("   devoluciones con pedido BC creado → comprobar abono en BC", f'=COUNTIF({rdv("Situación en BC")},"Pedido BC creado: comprobar abono/devolución en BC")', None, None, None, ""),
        ("   devoluciones con error BC (sin pedido)", f'=COUNTIF({rdv("Situación en BC")},"Error BC: sin pedido en BC")', None, None, None, ""),
        ("   devoluciones pendientes en BC (sin pedido ni abono)", f'=COUNTIF({rdv("Situación en BC")},"Pendiente en BC: sin pedido ni abono todavía")', None, None, None, ""),
        ("   devoluciones (abonos) creadas en BC", f'=COUNTIF({rdv("Nº devolución BC")},"?*")', None, None, None, "Líneas con nº de devolución BC en el log."),
        ("Líneas con Return Status en el log BC (REFUNDED / INCIDENT_OPEN)", f'=COUNTIF({rcl("Estado devolución (log BC)")},"REFUNDED")+COUNTIF({rcl("Estado devolución (log BC)")},"INCIDENT_OPEN")', None, None, None,
         "Lo que el log sabía en su fecha de extracción; el export es más reciente."),
    ]
    for etiqueta, n, np_, imp, reemb, nota in devk:
        _celda(ws, r, 2, etiqueta)
        _celda(ws, r, 3, n, "#,##0")
        _celda(ws, r, 4, np_, "#,##0")
        _celda(ws, r, 5, imp, FMT_EUR)
        _celda(ws, r, 6, reemb, FMT_EUR)
        _celda(ws, r, 7, nota, font=NOTA, borde=False)
        r += 1
    r += 1

    ws.cell(row=r, column=2, value="Contexto: log del conector completo, líneas por mes de creación y estado de proceso (valores calculados al generar el fichero)").font = BOLD
    r += 1
    estados = [c for c in ["Processed", "Error", "Pending"] if c in pivot_log.columns] + [c for c in pivot_log.columns if c not in ("Processed", "Error", "Pending")]
    _tabla_cabecera(ws, r, ["", "Mes de creación"] + estados + ["Total"] + [""] * max(0, 4 - len(estados)))
    r += 1
    for mes_txt, fila in pivot_log.iterrows():
        _celda(ws, r, 2, mes_txt)
        tot = 0
        for j, est in enumerate(estados):
            v = int(fila.get(est, 0))
            tot += v
            _celda(ws, r, 3 + j, v, "#,##0")
        _celda(ws, r, 3 + len(estados), tot, "#,##0", BOLD)
        r += 1
    r += 1
    ws.cell(row=r, column=2, value="Leyenda de la columna Devolución").font = BOLD
    r += 1
    for texto, rgb in COLOR_DEV.items():
        c = _celda(ws, r, 2, texto, font=BOLD)
        c.fill = PatternFill("solid", fgColor=rgb)
        _celda(ws, r, 7, {"Sí": "Línea reembolsada en Mirakl (estado Reembolsado o importe reembolsado > 0).",
                          "Incidencia": "Devolución en curso (estado Incidencia abierta, motivo Devolución)."}[texto], font=NOTA, borde=False)
        r += 1
    ws.sheet_view.showGridLines = False

    wb.calculation.fullCalcOnLoad = True
    salida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(salida)
    _sin_avisos_numero_texto(salida)
    return {"lineas_cruce": len(union), "pedidos": len(pedidos), "devoluciones": len(devs), "log": len(log), "log_periodo": len(la),
            "export": len(ex), "export_periodo": len(ea)}


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--bc", required=True, help="Excel 'Mirakl Orders Import' (log del conector BC)")
    ap.add_argument("--mirakl", required=True, help="Excel 'pedidos <año>' exportado de Mirakl (hoja 'orders')")
    ap.add_argument("--periodo", default=None, help="Mes a cuadrar, AAAA-MM (por defecto, el último mes completo del log)")
    ap.add_argument("--salida", default=None, help="Excel de salida (por defecto Cuadre_ECI_vs_BC_<periodo>.xlsx junto al log)")
    a = ap.parse_args()
    ruta_log, ruta_ex = Path(a.bc), Path(a.mirakl)
    print("Leyendo log BC…")
    log = leer_log_bc(ruta_log)
    print("Leyendo export Mirakl (puede tardar un minuto)…")
    ex = leer_export_mirakl(ruta_ex)
    if a.periodo:
        periodo = a.periodo
    else:
        ult = log["Fecha creación"].max()
        prev = (ult.replace(day=1) - dt.timedelta(days=1))
        periodo = prev.strftime("%Y-%m")
    salida = Path(a.salida) if a.salida else ruta_log.with_name(f"Cuadre_ECI_vs_BC_{periodo}.xlsx")
    info = construir(log, ex, periodo, salida, ruta_log, ruta_ex)
    print(f"Log BC: {info['log']} líneas ({info['log_periodo']} en {periodo}) | Export Mirakl: {info['export']} líneas ({info['export_periodo']} en {periodo}) | "
          f"Cruce: {info['lineas_cruce']} líneas, {info['pedidos']} pedidos | Devoluciones: {info['devoluciones']}")
    print(f"Guardado: {salida}")


if __name__ == "__main__":
    main()
