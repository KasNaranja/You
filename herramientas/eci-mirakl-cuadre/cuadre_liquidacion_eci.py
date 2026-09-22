# -*- coding: utf-8 -*-
"""
Cuadre de las liquidaciones de El Corte Inglés («Liquidación explotaciones directas») con los
pedidos del marketplace (export de Mirakl).

Regla de ECI (deducida y comprobada con las liquidaciones de agosto de 2026):
  · La liquidación va por centro y por departamento (UNECO): el centro está en el propio número
    de pedido de Mirakl (posiciones 4 a 7: 0090 = venta a distancia España, 0143 = Portugal,
    0005 = Bilbao…) y el departamento sale del producto (mujer/unisex → 0601, hombre → 0602,
    niños → 0696).
  · VENTA BRUTA del mes = suma del importe de las líneas con «Fecha de cumplimentación» en el mes,
    menos el importe reembolsado de las líneas con «Fecha de abono 1» en el mes que llegaron a
    cumplimentarse (los reembolsos de líneas nunca entregadas no restan porque nunca sumaron).

Entradas
  --mirakl         'pedidos <año>.xlsx' exportado de Mirakl (hoja 'orders')
  --liquidaciones  carpeta con las liquidaciones del mes: los .xlsx de ECI España (hoja 'Data'),
                   los .XLS de Portugal (texto UTF-16 con tabuladores), la subcarpeta 'corners'
                   y, si existe, el '...total....xlsx' del usuario (para los nº de factura)
  --periodo        AAAA-MM (por defecto, el mes de FE.CONTAB de las liquidaciones)
Salida
  --salida  Excel con Resumen (liquidación vs Mirakl por línea de liquidación), Líneas mes
            (todas las líneas de Mirakl con actividad en el mes, con su venta y abono en
            liquidación por fórmula) y Liquidación ECI (las líneas tal cual vienen de ECI).
"""
import argparse
import calendar
import datetime as dt
import glob
import os
import re
import shutil
import sys
import warnings
import zipfile
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore")

FUENTE = "Arial"
MAXF = 8000
CAB_FILL = PatternFill("solid", fgColor="1F4E78")
CAB_FONT = Font(name=FUENTE, size=10, bold=True, color="FFFFFF")
NORMAL = Font(name=FUENTE, size=10)
BOLD = Font(name=FUENTE, size=10, bold=True)
TITULO = Font(name=FUENTE, size=14, bold=True, color="1F4E78")
NOTA = Font(name=FUENTE, size=9, italic=True, color="595959")
AZUL = Font(name=FUENTE, size=10, color="0000FF")
THIN = Side(style="thin", color="BFBFBF")
BORDE = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FMT_FECHA = "dd/mm/yyyy hh:mm"
FMT_DIA = "dd/mm/yyyy"
FMT_EUR = '#,##0.00 "€"'
FMT_DIF = '#,##0.00 "€";[Red]-#,##0.00 "€";"-"'
FMT_PCT = "0.00%"
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
         "septiembre", "octubre", "noviembre", "diciembre"]
DEPTOS = {"0601": "Zapatería dig. mujer", "0602": "Zapatería dig. hombre", "0696": "Zapatería dig. niños",
          "0598": "Zapatería E.D. (corners, venta física)"}
COLOR_EST = {"Cuadra": "C6EFCE", "Diferencia": "FFEB9C", "Solo en liquidación": "FFC7CE", "Solo en Mirakl": "F4B183",
             "Corners: fuera del marketplace": "E7E6E6"}


# ----------------------------------------------------------------------------- lectura
def _num(v):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("'", "").replace(" ", "")
    if s in ("", "-"):
        return 0.0
    neg = s.endswith("-")
    s = s.rstrip("-")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        x = float(s)
    except ValueError:
        return 0.0
    return -x if neg else x


def _centro(v):
    s = str(v).strip().replace("'", "")
    m = re.fullmatch(r"E(\d{3})E", s)
    if m:
        return "0" + m.group(1)
    if s.isdigit():
        return s[-4:].zfill(4)
    return s


def _uneco(v):
    s = str(v).strip().replace("'", "")
    return s.zfill(4) if s.isdigit() else s


def _fecha(v):
    if isinstance(v, (dt.datetime, dt.date, pd.Timestamp)):
        return pd.Timestamp(v)
    s = str(v).strip().replace("'", "")
    for f in ("%d.%m.%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return pd.Timestamp(dt.datetime.strptime(s, f))
        except ValueError:
            pass
    return pd.NaT


COLS_LIQ = ["CENTRO", "DESCR.CENTRO", "FE.CONTAB", "UNECO", "NoPEDIDO", "NoALBARAN", "TOTAL VENTA", "NETA VENTA", "IVA VENTA",
            "%PARTICIP.ECI", "EUROS PARTICIP.ECI", "IMPORTE A FACTURAR POR PROV", "IVA A FACTURAR POR PROV", "TOTAL FACTURA"]


def _filas_liq(cab, filas, pais, fichero, tipo):
    """Normaliza filas de una liquidación de ECI (mismas 14 columnas en ES y PT)."""
    cab = [str(c).strip().rstrip(".") for c in cab]
    idx = {c: i for i, c in enumerate(cab)}
    out = []
    for r in filas:
        if r is None or all(v is None or str(v).strip() == "" for v in r):
            continue
        g = lambda c: r[idx[c]] if c in idx and idx[c] < len(r) else None
        if g("CENTRO") is None or g("TOTAL VENTA") is None:
            continue
        out.append({
            "País": pais, "Fichero": fichero, "Tipo": tipo,
            "Centro": _centro(g("CENTRO")), "Descripción centro": str(g("DESCR.CENTRO") or "").strip(),
            "Fecha contable": _fecha(g("FE.CONTAB")), "UNECO": _uneco(g("UNECO")),
            "Nº pedido ECI": str(g("NoPEDIDO") or "").replace("'", "").strip(), "Nº albarán ECI": str(g("NoALBARAN") or "").replace("'", "").strip(),
            "Total venta (bruta)": _num(g("TOTAL VENTA")), "Neta venta": _num(g("NETA VENTA")), "IVA venta": _num(g("IVA VENTA")),
            "% particip. ECI": _num(g("%PARTICIP.ECI")) / 100.0, "Particip. ECI €": _num(g("EUROS PARTICIP.ECI")),
            "Importe a facturar": _num(g("IMPORTE A FACTURAR POR PROV")), "IVA a facturar": _num(g("IVA A FACTURAR POR PROV")),
            "Total factura": _num(g("TOTAL FACTURA")),
        })
    return out


def leer_liquidaciones(carpeta: Path):
    filas, otros = [], []
    for ruta in sorted(carpeta.rglob("*")):
        if not ruta.is_file() or ruta.name.startswith("~$"):
            continue
        rel = ruta.relative_to(carpeta).as_posix()
        tipo = "Corners (venta física)" if "corner" in rel.lower() else "Marketplace"
        low = ruta.name.lower()
        if ruta.suffix.lower() == ".xlsx" and "total" not in low and "abono" not in low and "cuadre" not in low:
            wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
            if "Data" in wb.sheetnames:
                ws = wb["Data"]
                rows = list(ws.iter_rows(values_only=True))
                if rows:
                    filas += _filas_liq(rows[0], rows[1:], "ES", rel, tipo)
            else:
                otros.append(rel)
            wb.close()
        elif ruta.suffix.lower() == ".xls":
            try:
                txt = ruta.read_bytes().decode("utf-16")
            except UnicodeDecodeError:
                otros.append(rel)
                continue
            lineas = [l for l in re.split(r"\r\n|\r|\n", txt) if l.strip()]
            rows = [l.split("\t") for l in lineas]
            if rows and "CENTRO" in rows[0][0]:
                filas += _filas_liq(rows[0], rows[1:], "PT", rel, tipo)
            else:
                otros.append(rel)
        elif ruta.suffix.lower() in (".pdf", ".xlsx"):
            otros.append(rel)
    liq = pd.DataFrame(filas)
    if liq.empty:
        sys.exit("No he encontrado ninguna liquidación en la carpeta")
    liq["Departamento"] = liq["UNECO"].map(lambda u: DEPTOS.get(u, u))
    liq["Negativa"] = liq["Total venta (bruta)"] < 0
    liq = liq.drop_duplicates(subset=["País", "UNECO", "Centro", "Nº albarán ECI", "Total venta (bruta)"]).reset_index(drop=True)
    # nº de factura del fichero consolidado del usuario, si existe
    fras = {}
    for ruta in carpeta.glob("*total*.xlsx"):
        wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        cab = [str(c).strip() if c is not None else "" for c in rows[0]]
        if "NoALBARAN" in cab and "nº fra SAP" in cab:
            ia, ifr = cab.index("NoALBARAN"), cab.index("nº fra SAP")
            for r in rows[1:]:
                if r and r[ia] and r[ifr]:
                    fras[str(r[ia]).replace("'", "").strip()] = str(r[ifr]).strip()
        wb.close()
    liq["Nº factura (fichero total)"] = liq["Nº albarán ECI"].map(fras).fillna("")
    return liq, otros


def leer_export_mirakl(ruta: Path) -> pd.DataFrame:
    df = pd.read_excel(ruta, sheet_name="orders", dtype=str, keep_default_na=False)
    df.columns = [str(c).strip() for c in df.columns]
    df = df[df["N.º de asiento de pedido"].str.strip() != ""].copy()
    df["Fecha creación"] = pd.to_datetime(df["Fecha de creación"], format="%d/%m/%Y %H:%M:%S", errors="coerce")
    for c in ["Fecha de envío", "Entrega", "Fecha de recepción"]:
        df[c] = pd.to_datetime(df[c].replace("", None), format="%d/%m/%Y %H:%M:%S", errors="coerce")
    for c in ["Fecha Emisión", "Fecha de cumplimentación", "Fecha de abono 1", "Fecha de abono 2", "Fecha de recogida 1", "Fecha de recogida 2"]:
        df[c] = pd.to_datetime(df[c].replace("", None), format="%Y%m%d", errors="coerce")
    for c in ["Cantidad", "Importe", "Importe total reembolsado (impuestos incluidos)", "Importe total cancelado (impuestos incluidos)",
              "Comisión (sin impuestos)", "Importe transferido a tienda (impuestos incluidos)"]:
        df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0)
    df["País"] = df["Canal"].map(lambda c: "PT" if "Portugal" in str(c) else "ES")
    df["Centro"] = df["Número de pedido"].str[3:7]
    df["UNECO"] = df["Detalles"].map(departamento)
    return df.reset_index(drop=True)


def departamento(detalles) -> str:
    t = str(detalles).lower()
    if re.search(r"niñ|infantil|beb[eé]|junior|kids|\bchic[oa]s?\b|crian[cç]a|menin[oa]|rapaz|rapariga", t):
        return "0696"
    if re.search(r"hombre|homem|\bmen\b", t):
        return "0602"
    return "0601"


# ----------------------------------------------------------------------------- escritura
def _cabecera(ws, fila, cabeceras, anchos=None):
    for j, h in enumerate(cabeceras, start=1):
        c = ws.cell(row=fila, column=j, value=h)
        c.fill, c.font, c.border = CAB_FILL, CAB_FONT, BORDE
        c.alignment = Alignment(vertical="center", wrap_text=True)
        if anchos and j - 1 < len(anchos) and anchos[j - 1]:
            ws.column_dimensions[get_column_letter(j)].width = anchos[j - 1]
    ws.row_dimensions[fila].height = 30


def _celda(ws, fila, col, valor, fmt=None, font=NORMAL, borde=True):
    if isinstance(valor, float) and pd.isna(valor):
        valor = None
    elif isinstance(valor, pd.Timestamp):
        valor = None if pd.isna(valor) else valor.to_pydatetime()
    elif valor is pd.NaT:
        valor = None
    c = ws.cell(row=fila, column=col, value=valor)
    c.font = font
    if borde:
        c.border = BORDE
    if fmt:
        c.number_format = fmt
    return c


def _pinta_filas(ws, col_clave, primera, ultima, ncols, colores):
    rango = f"A{primera}:{get_column_letter(ncols)}{ultima}"
    letra = get_column_letter(col_clave)
    for texto, rgb in colores.items():
        ws.conditional_formatting.add(rango, FormulaRule(formula=[f'${letra}{primera}="{texto}"'], fill=PatternFill("solid", fgColor=rgb, bgColor=rgb)))


def _escribe_df(ws, df, fila_cab, anchos, formatos):
    _cabecera(ws, fila_cab, list(df.columns), anchos)
    cols = list(df.columns)
    for i, fila in enumerate(df.itertuples(index=False), start=fila_cab + 1):
        for j, v in enumerate(fila, start=1):
            _celda(ws, i, j, v, formatos.get(cols[j - 1]))
    ws.freeze_panes = ws.cell(row=fila_cab + 1, column=2)
    ws.auto_filter.ref = f"A{fila_cab}:{get_column_letter(len(cols))}{max(fila_cab + 1, fila_cab + len(df))}"


def _sin_avisos_numero_texto(ruta: Path):
    tag = f'<ignoredErrors><ignoredError sqref="A1:BZ{MAXF}" numberStoredAsText="1"/></ignoredErrors>'
    tmp = ruta.with_suffix(".tmp.xlsx")
    with zipfile.ZipFile(ruta) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            datos = zin.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet") and item.filename.endswith(".xml"):
                xml = datos.decode("utf-8")
                if "<ignoredErrors" not in xml and "</worksheet>" in xml and "<tableParts" not in xml and "<drawing" not in xml:
                    datos = xml.replace("</worksheet>", tag + "</worksheet>").encode("utf-8")
            zout.writestr(item, datos)
    shutil.move(tmp, ruta)


def construir(ex, liq, otros, periodo, salida, ruta_ex, carpeta):
    anyo, mes = (int(x) for x in periodo.split("-"))
    ini = dt.datetime(anyo, mes, 1)
    fin = dt.datetime(anyo, mes, calendar.monthrange(anyo, mes)[1], 23, 59, 59)
    nombre_mes = MESES[mes - 1]
    en_mes = lambda s: (s >= ini) & (s <= fin)
    cumpl = ex["Fecha de cumplimentación"].notna()
    act = ex[en_mes(ex["Fecha de cumplimentación"]) | en_mes(ex["Fecha de abono 1"]) | en_mes(ex["Fecha de abono 2"]) | en_mes(ex["Fecha de recogida 1"])].copy()
    act = act.sort_values(["País", "Centro", "UNECO", "Fecha de cumplimentación", "Fecha creación"]).reset_index(drop=True)
    nombres = dict(zip(liq["Centro"], liq["Descripción centro"]))
    # marcas de revisión (datos, calculadas al generar)
    def marca(r):
        m = []
        if r["Estado"] in ("Rechazado", "Cancelado") and pd.notna(r["Fecha de cumplimentación"]):
            m.append("Rechazada/cancelada con cumplimentación")
        if pd.notna(r["Fecha de abono 1"]) and en_mes(pd.Series([r["Fecha de abono 1"]])).iloc[0] and r["Importe total reembolsado (impuestos incluidos)"] == 0 and r["Estado"] not in ("Rechazado", "Cancelado"):
            m.append("Talón de abono sin reembolso en Mirakl")
        if pd.notna(r["Fecha de abono 1"]) and en_mes(pd.Series([r["Fecha de abono 1"]])).iloc[0] and r["Importe total reembolsado (impuestos incluidos)"] > 0 and pd.isna(r["Fecha de cumplimentación"]):
            m.append("Reembolso de línea no cumplimentada (no resta)")
        if pd.notna(r["Fecha de recogida 1"]) and en_mes(pd.Series([r["Fecha de recogida 1"]])).iloc[0] and pd.notna(r["Fecha de abono 1"]) and r["Fecha de abono 1"] > fin and r["Importe total reembolsado (impuestos incluidos)"] > 0:
            m.append("Devolución recogida en el mes y abonada después")
        if pd.notna(r["Fecha de cumplimentación"]) and r["Fecha de cumplimentación"].date() == fin.date():
            m.append("Cumplimentada el último día del mes")
        if 0 < r["Importe total reembolsado (impuestos incluidos)"] < r["Importe"] - 0.005:
            m.append("Reembolso parcial")
        return " | ".join(m)
    act["Revisar"] = act.apply(marca, axis=1)

    wb = openpyxl.Workbook()
    ws_res = wb.active
    ws_res.title = "Resumen"
    ws_lin = wb.create_sheet("Líneas mes")
    ws_liq = wb.create_sheet("Liquidación ECI")

    # --- hoja Liquidación ECI -------------------------------------------------------
    cols_liq = ["País", "Tipo", "UNECO", "Departamento", "Centro", "Descripción centro", "Fecha contable", "Nº pedido ECI", "Nº albarán ECI",
                "Nº factura (fichero total)", "Total venta (bruta)", "Neta venta", "IVA venta", "% particip. ECI", "Particip. ECI €",
                "Importe a facturar", "IVA a facturar", "Total factura", "Fichero"]
    ws_liq["A1"] = f"Liquidaciones de ECI leídas de la carpeta «{carpeta.name}» ({len(liq)} líneas)"
    ws_liq["A1"].font = TITULO
    ws_liq["A2"] = "Tal cual vienen en los ficheros de ECI (España: hoja Data de cada .xlsx; Portugal: .XLS de texto). 'Nº factura' sale del fichero «total» del usuario, si existe."
    ws_liq["A2"].font = NOTA
    FL = 4
    _escribe_df(ws_liq, liq[cols_liq], FL, [6, 22, 8, 26, 8, 32, 12, 12, 12, 14, 14, 12, 11, 9, 12, 13, 12, 13, 44],
                {"Fecha contable": FMT_DIA, "Total venta (bruta)": FMT_EUR, "Neta venta": FMT_EUR, "IVA venta": FMT_EUR, "% particip. ECI": FMT_PCT,
                 "Particip. ECI €": FMT_EUR, "Importe a facturar": FMT_EUR, "IVA a facturar": FMT_EUR, "Total factura": FMT_EUR})
    LQ = {c: get_column_letter(i + 1) for i, c in enumerate(cols_liq)}

    def rlq(col):
        return f"'Liquidación ECI'!${LQ[col]}${FL + 1}:${LQ[col]}${MAXF}"

    # --- hoja Líneas mes -------------------------------------------------------------
    cab = ["N.º asiento", "Nº pedido Mirakl", "País", "Centro", "Nombre centro", "UNECO", "Departamento", "Estado Mirakl", "Motivo",
           "SKU (EAN)", "Descripción", "Cantidad", "Importe", "Importe reembolsado", "Fecha creación", "Fecha cumplimentación",
           "Talón cumplimentación", "Fecha abono 1", "Talón abono 1", "Fecha abono 2", "Fecha recogida 1", "Talón recogida 1", "Método de envío",
           "Venta en liquidación (mes)", "Abono en liquidación (mes)", "Neto liquidación", "Revisar"]
    ws_lin["A1"] = f"Líneas de Mirakl con actividad en {nombre_mes} de {anyo} (cumplimentación, abono o recogida en el mes)"
    ws_lin["A1"].font = TITULO
    ws_lin["A2"] = ("«Venta en liquidación» = importe si la fecha de cumplimentación cae en el periodo del Resumen; «Abono en liquidación» = importe reembolsado "
                    "si la fecha de abono 1 cae en el periodo y la línea tiene fecha de cumplimentación. Centro = posiciones 4-7 del nº de pedido; UNECO por el producto.")
    ws_lin["A2"].font = NOTA
    FC = 4
    _cabecera(ws_lin, FC, cab, [34, 32, 6, 8, 26, 8, 22, 16, 12, 15, 40, 8, 11, 12, 16, 14, 12, 12, 12, 12, 12, 12, 22, 13, 13, 13, 50])
    L = {c: get_column_letter(i + 1) for i, c in enumerate(cab)}
    p1 = FC + 1
    col_src = {"N.º asiento": "N.º de asiento de pedido", "Nº pedido Mirakl": "Número de pedido", "País": "País", "Centro": "Centro", "UNECO": "UNECO",
               "Estado Mirakl": "Estado", "Motivo": "Motivo", "SKU (EAN)": "SKU del producto", "Descripción": "Detalles", "Cantidad": "Cantidad",
               "Importe": "Importe", "Importe reembolsado": "Importe total reembolsado (impuestos incluidos)", "Fecha creación": "Fecha creación",
               "Fecha cumplimentación": "Fecha de cumplimentación", "Talón cumplimentación": "Talón de cumplimentación", "Fecha abono 1": "Fecha de abono 1",
               "Talón abono 1": "Talón de abono 1", "Fecha abono 2": "Fecha de abono 2", "Fecha recogida 1": "Fecha de recogida 1",
               "Talón recogida 1": "Talón de recogida 1", "Método de envío": "Método de envío", "Revisar": "Revisar"}
    fmts = {"Importe": FMT_EUR, "Importe reembolsado": FMT_EUR, "Fecha creación": FMT_FECHA, "Fecha cumplimentación": FMT_DIA, "Fecha abono 1": FMT_DIA,
            "Fecha abono 2": FMT_DIA, "Fecha recogida 1": FMT_DIA, "Venta en liquidación (mes)": FMT_EUR, "Abono en liquidación (mes)": FMT_EUR,
            "Neto liquidación": FMT_DIF, "Cantidad": "0"}
    for i, (_, r) in enumerate(act.iterrows(), start=p1):
        for col in cab:
            if col in col_src:
                v = r[col_src[col]]
                if col == "Descripción":
                    v = str(v)[:80]
                _celda(ws_lin, i, cab.index(col) + 1, v, fmts.get(col))
        _celda(ws_lin, i, cab.index("Nombre centro") + 1, nombres.get(r["Centro"], ""))
        _celda(ws_lin, i, cab.index("Departamento") + 1, DEPTOS.get(r["UNECO"], r["UNECO"]))
        fc, fa, imp, reemb = f'{L["Fecha cumplimentación"]}{i}', f'{L["Fecha abono 1"]}{i}', f'{L["Importe"]}{i}', f'{L["Importe reembolsado"]}{i}'
        _celda(ws_lin, i, cab.index("Venta en liquidación (mes)") + 1, f'=IF(AND({fc}<>"",{fc}>=Resumen!$C$5,{fc}<=Resumen!$C$6),{imp},0)', FMT_EUR)
        _celda(ws_lin, i, cab.index("Abono en liquidación (mes)") + 1, f'=IF(AND({fa}<>"",{fa}>=Resumen!$C$5,{fa}<=Resumen!$C$6,{fc}<>""),{reemb},0)', FMT_EUR)
        _celda(ws_lin, i, cab.index("Neto liquidación") + 1, f'={L["Venta en liquidación (mes)"]}{i}-{L["Abono en liquidación (mes)"]}{i}', FMT_DIF)
    pN = p1 + max(len(act), 1) - 1
    ws_lin.freeze_panes = f"B{p1}"
    ws_lin.auto_filter.ref = f"A{FC}:{get_column_letter(len(cab))}{pN}"
    cR = L["Revisar"]
    ws_lin.conditional_formatting.add(f"{cR}{p1}:{cR}{pN}", FormulaRule(formula=[f'${cR}{p1}<>""'], fill=PatternFill("solid", fgColor="FFF2CC", bgColor="FFF2CC")))

    def rln(col):
        return f"'Líneas mes'!${L[col]}${p1}:${L[col]}${pN}"

    # --- hoja Resumen ----------------------------------------------------------------
    ws = ws_res
    for col, w in zip("ABCDEFGHIJKLMN", [3, 6, 8, 24, 8, 30, 14, 14, 14, 14, 14, 22, 14, 60]):
        ws.column_dimensions[col].width = w
    ws["B1"] = f"Cuadre de la liquidación de ECI con los pedidos de Mirakl — {nombre_mes} {anyo}"
    ws["B1"].font = TITULO
    ws["B2"] = (f"Generado el {dt.date.today():%d/%m/%Y} a partir de las liquidaciones de «{carpeta.name}» y de '{ruta_ex.name}'. "
                "Regla: venta bruta del mes = líneas cumplimentadas en el mes − reembolsos abonados en el mes de líneas cumplimentadas, por centro y departamento.")
    ws["B2"].font = NOTA
    ws["B4"] = "Periodo (fechas de cumplimentación y de abono)"
    ws["B4"].font = BOLD
    _celda(ws, 5, 2, "Inicio")
    _celda(ws, 5, 3, ini, FMT_DIA, AZUL)
    _celda(ws, 6, 2, "Fin")
    _celda(ws, 6, 3, fin, FMT_FECHA, AZUL)
    r = 8
    ws.cell(row=r, column=2, value="Liquidación de ECI frente a Mirakl, línea a línea").font = BOLD
    r += 1
    cab_r = ["", "País", "UNECO", "Departamento", "Centro", "Descripción centro", "Total venta liquidación", "Ventas Mirakl (cumplimentadas)",
             "Abonos Mirakl (reembolsos)", "Neto Mirakl", "Diferencia (Mirakl − ECI)", "Estado", "Nº factura", "Notas"]
    _cabecera(ws, r, cab_r)
    ws.cell(row=r, column=1).fill = PatternFill(fill_type=None)
    ws.cell(row=r, column=1).border = Border()
    r += 1
    r_ini = r
    # filas: liquidaciones marketplace (agrupadas por país/UNECO/centro) + combinaciones Mirakl con neto ≠ 0 que no estén
    mk = liq[liq["Tipo"] == "Marketplace"].groupby(["País", "UNECO", "Centro"], as_index=False).agg(
        descr=("Descripción centro", "first"), total=("Total venta (bruta)", "sum"), fra=("Nº factura (fichero total)", lambda s: " / ".join(sorted(set(x for x in s if x)))))
    claves = set(zip(mk["País"], mk["UNECO"], mk["Centro"]))
    ventas = act[en_mes(act["Fecha de cumplimentación"])].groupby(["País", "UNECO", "Centro"])["Importe"].sum()
    abonos = act[en_mes(act["Fecha de abono 1"]) & act["Fecha de cumplimentación"].notna()].groupby(["País", "UNECO", "Centro"])["Importe total reembolsado (impuestos incluidos)"].sum()
    neto = (ventas.subtract(abonos, fill_value=0)).round(2)
    extra = [k for k, v in neto.items() if abs(v) >= 0.005 and k not in claves]
    filas = [(p, u, c, d, t, f, True) for p, u, c, d, t, f in zip(mk["País"], mk["UNECO"], mk["Centro"], mk["descr"], mk["total"], mk["fra"])]
    filas += [(p, u, c, nombres.get(c, ""), None, "", False) for (p, u, c) in extra]
    filas.sort(key=lambda t: (t[0], t[1], t[2]))
    for pais, uneco, centro, descr, total, fra, en_liq in filas:
        _celda(ws, r, 2, pais)
        _celda(ws, r, 3, uneco)
        _celda(ws, r, 4, DEPTOS.get(uneco, uneco))
        _celda(ws, r, 5, centro)
        _celda(ws, r, 6, descr)
        _celda(ws, r, 7, f'=SUMIFS({rlq("Total venta (bruta)")},{rlq("País")},$B{r},{rlq("UNECO")},$C{r},{rlq("Centro")},$E{r},{rlq("Tipo")},"Marketplace")' if en_liq else None, FMT_EUR)
        _celda(ws, r, 8, f'=SUMIFS({rln("Venta en liquidación (mes)")},{rln("País")},$B{r},{rln("UNECO")},$C{r},{rln("Centro")},$E{r})', FMT_EUR)
        _celda(ws, r, 9, f'=SUMIFS({rln("Abono en liquidación (mes)")},{rln("País")},$B{r},{rln("UNECO")},$C{r},{rln("Centro")},$E{r})', FMT_EUR)
        _celda(ws, r, 10, f"=H{r}-I{r}", FMT_EUR)
        _celda(ws, r, 11, f"=ROUND(J{r}-N(G{r}),2)", FMT_DIF)
        _celda(ws, r, 12, (f'=IF(G{r}="","Solo en Mirakl",IF(ABS(K{r})<0.01,"Cuadra","Diferencia"))' if en_liq else '="Solo en Mirakl"'))
        _celda(ws, r, 13, fra)
        _celda(ws, r, 14, None, font=NOTA)
        r += 1
    r_fin = r - 1
    _celda(ws, r, 6, "Total marketplace", font=BOLD)
    for col in "GHIJK":
        _celda(ws, r, ord(col) - 64, f"=SUM({col}{r_ini}:{col}{r_fin})", FMT_EUR if col != "K" else FMT_DIF, BOLD)
    _pinta_filas(ws, 12, r_ini, r_fin, 14, COLOR_EST)
    r += 2
    ws.cell(row=r, column=2, value="Totales por país y departamento").font = BOLD
    r += 1
    _cabecera(ws, r, ["", "País", "UNECO", "Departamento", "", "", "Total venta liquidación", "Ventas Mirakl", "Abonos Mirakl", "Neto Mirakl", "Diferencia", "Desviación %"])
    ws.cell(row=r, column=1).fill = PatternFill(fill_type=None)
    ws.cell(row=r, column=1).border = Border()
    r += 1
    r_t = r
    for pais, uneco in sorted(set((p, u) for p, u, *_ in filas)):
        _celda(ws, r, 2, pais)
        _celda(ws, r, 3, uneco)
        _celda(ws, r, 4, DEPTOS.get(uneco, uneco))
        for col in "GHIJK":
            _celda(ws, r, ord(col) - 64, f'=SUMIFS({col}{r_ini}:{col}{r_fin},$B${r_ini}:$B${r_fin},$B{r},$C${r_ini}:$C${r_fin},$C{r})', FMT_EUR if col != "K" else FMT_DIF)
        _celda(ws, r, 12, f'=IF(G{r}=0,"",K{r}/G{r})', FMT_PCT)
        r += 1
    _celda(ws, r, 4, "Total", font=BOLD)
    for col in "GHIJK":
        _celda(ws, r, ord(col) - 64, f"=SUM({col}{r_t}:{col}{r - 1})", FMT_EUR if col != "K" else FMT_DIF, BOLD)
    _celda(ws, r, 12, f'=IF(G{r}=0,"",K{r}/G{r})', FMT_PCT, BOLD)
    r += 2
    ws.cell(row=r, column=2, value="Posibles causas de las diferencias (importes de las líneas marcadas en 'Líneas mes', columna Revisar)").font = BOLD
    r += 1
    _cabecera(ws, r, ["", "Marca", "", "", "", "", "Líneas", "Importe líneas", "Importe reembolsado", "", "", "", "", "Qué significa"])
    ws.cell(row=r, column=1).fill = PatternFill(fill_type=None)
    ws.cell(row=r, column=1).border = Border()
    r += 1
    for marca_txt, expl in [
        ("Rechazada/cancelada con cumplimentación", "Líneas rechazadas o canceladas que sin embargo tienen talón de cumplimentación: se han sumado como venta; ECI puede no contarlas o abonarlas."),
        ("Talón de abono sin reembolso en Mirakl", "ECI tiene talón de abono en el mes pero Mirakl no muestra reembolso (estado Recibido). No se han restado."),
        ("Reembolso de línea no cumplimentada (no resta)", "Reembolsos de líneas que nunca se entregaron: no se restan porque tampoco sumaron. Confirmado con León, Costa Luz, Marbella y A Coruña."),
        ("Devolución recogida en el mes y abonada después", "La devolución se recogió en el mes pero el abono de Mirakl es posterior: ECI puede haberla restado ya (caso Valderas)."),
        ("Cumplimentada el último día del mes", "Líneas del último día: posible corte distinto en ECI."),
        ("Reembolso parcial", "Reembolso menor que el importe de la línea (2 unidades, 1 devuelta)."),
    ]:
        _celda(ws, r, 2, marca_txt)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        _celda(ws, r, 7, f'=COUNTIF({rln("Revisar")},"*{marca_txt}*")', "#,##0")
        _celda(ws, r, 8, f'=SUMIF({rln("Revisar")},"*{marca_txt}*",{rln("Importe")})', FMT_EUR)
        _celda(ws, r, 9, f'=SUMIF({rln("Revisar")},"*{marca_txt}*",{rln("Importe reembolsado")})', FMT_EUR)
        _celda(ws, r, 14, expl, font=NOTA, borde=False)
        r += 1
    r += 1
    ws.cell(row=r, column=2, value="Otras líneas y documentos de la carpeta que no se cruzan con Mirakl").font = BOLD
    r += 1
    corners = liq[liq["Tipo"] != "Marketplace"]
    for _, c in corners.iterrows():
        _celda(ws, r, 2, f"{c['Tipo']}: {c['Centro']} {c['Descripción centro']} (UNECO {c['UNECO']}, {c['% particip. ECI']:.0%} participación) → total venta {c['Total venta (bruta)']:,.2f} €, a facturar {c['Total factura']:,.2f} €".replace(",", "X").replace(".", ",").replace("X", "."), font=NORMAL, borde=False)
        r += 1
    for o in otros:
        _celda(ws, r, 2, f"Documento no cruzado: {o}", font=NOTA, borde=False)
        r += 1
    r += 1
    _celda(ws, r, 2, "Leyenda", font=BOLD, borde=False)
    r += 1
    for texto, rgb in COLOR_EST.items():
        c = _celda(ws, r, 2, texto)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        c.fill = PatternFill("solid", fgColor=rgb)
        r += 1
    ws.sheet_view.showGridLines = False

    wb.calculation.fullCalcOnLoad = True
    wb.save(salida)
    _sin_avisos_numero_texto(salida)
    return {"liq": len(liq), "lineas": len(act), "filas_resumen": len(filas), "extra": len(extra)}


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--mirakl", required=True)
    ap.add_argument("--liquidaciones", required=True, help="carpeta con las liquidaciones del mes")
    ap.add_argument("--periodo", default=None, help="AAAA-MM")
    ap.add_argument("--salida", default=None)
    a = ap.parse_args()
    carpeta = Path(a.liquidaciones)
    liq, otros = leer_liquidaciones(carpeta)
    print(f"Liquidaciones: {len(liq)} líneas ({liq['País'].value_counts().to_dict()}), otros documentos: {len(otros)}")
    print("Leyendo export Mirakl (puede tardar un minuto)…")
    ex = leer_export_mirakl(Path(a.mirakl))
    periodo = a.periodo or liq["Fecha contable"].dropna().max().strftime("%Y-%m")
    salida = Path(a.salida) if a.salida else carpeta / f"Cuadre_liquidacion_ECI_{periodo}.xlsx"
    info = construir(ex, liq, otros, periodo, salida, Path(a.mirakl), carpeta)
    print(f"Periodo {periodo} | líneas Mirakl con actividad: {info['lineas']} | filas del resumen: {info['filas_resumen']} (solo Mirakl: {info['extra']})")
    print(f"Guardado: {salida}")


if __name__ == "__main__":
    main()
