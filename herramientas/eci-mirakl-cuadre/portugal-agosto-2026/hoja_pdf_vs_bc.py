# -*- coding: utf-8 -*-
"""Hoja 'PDF vs BC': cada columna de los PDF de liquidación de ECI Portugal frente al campo equivalente de BC."""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

F = 'Arial'
CAB = PatternFill('solid', fgColor='1F4E78')
SUB = PatternFill('solid', fgColor='D9E1F2')
CABF = Font(name=F, size=10, bold=True, color='FFFFFF')
N = Font(name=F, size=10)
B = Font(name=F, size=10, bold=True)
AZ = Font(name=F, size=10, color='0000FF')
T = Font(name=F, size=14, bold=True, color='1F4E78')
NOTA = Font(name=F, size=9, italic=True, color='595959')
th = Side(style='thin', color='BFBFBF')
BOR = Border(left=th, right=th, top=th, bottom=th)
EUR = '#,##0.00 "€";-#,##0.00 "€";"-"'
DIF = '#,##0.00 "€";[Red]-#,##0.00 "€";"-"'
PCT = '0.00%'

# Datos de los PDF de ECI Portugal (LIQUIDAÇÃO EXPLORAÇÕES DIRECTAS_1/2/3, agosto 2026)
PDF = {
    '601': dict(total=4922.72, neta=4002.21, iva=920.51, pct=0.29, part=1160.64, fact=2841.57, ivaf=0.0, totf=2841.57),
    '602': dict(total=981.56, neta=798.02, iva=183.54, pct=0.29, part=231.43, fact=566.59, ivaf=0.0, totf=566.59),
    '696': dict(total=101.87, neta=82.82, iva=19.05, pct=0.29, part=24.02, fact=58.80, ivaf=0.0, totf=58.80),
}
BC_COLS = ['Importe liquidable', 'Base imponible comisión', 'Comisión esperada', 'IVA comisión esperado', 'Importe neto esperado']

# (concepto, nombre en el PDF, campo BC mostrado, clave PDF, campo BC para la fórmula, comentario)
FILAS = [
    ('Venta total con IVA', 'VENDA TOTAL', 'Importe liquidable', 'total', 'Importe liquidable',
     'Igual. ECI suma en 0143 las devoluciones hechas en las tiendas de Lisboa (0140) y Gaia (0142).'),
    ('Venta neta sin IVA', 'NETA', 'Base imponible comisión', 'neta', 'Base imponible comisión',
     'BC no quita el IVA: su base de comisión es el importe con IVA (y sin las líneas sin vincular).'),
    ('IVA de la venta (23 %)', 'IVA (23.00)', '(no existe en BC)', 'iva', None,
     'BC no separa el IVA de la venta en este informe.'),
    ('% participación ECI', '%', 'Comisión esperada / base', 'pct', 'PCT',
     'El porcentaje es el mismo (29 %); cambia la base sobre la que se aplica.'),
    ('Participación ECI (comisión)', 'PARTICIP.ECI EUROS', 'Comisión esperada', 'part', 'Comisión esperada',
     'ECI: 29 % de la neta sin IVA. BC: 29 % del importe con IVA.'),
    ('Importe a facturar (base)', 'A FACTURAR IMPORTE', 'Importe neto esperado', 'fact', 'Importe neto esperado',
     'ECI: neta − participación. BC: importe con IVA − comisión.'),
    ('IVA a facturar (0 %)', 'IVA (0.00)', 'IVA comisión esperado', 'ivaf', 'IVA comisión esperado',
     'Portugal se factura sin IVA. En BC este campo es el IVA de la comisión, otro concepto; vale 0.'),
    ('Total factura', 'TOTAL FRA', 'Importe neto esperado', 'totf', 'Importe neto esperado',
     'BC no tiene un total de factura; lo más parecido es el neto esperado.'),
]


def _c(ws, r, c, v, font=N, fmt=None, fill=None, wrap=False, center=False):
    x = ws.cell(row=r, column=c, value=v)
    x.font, x.border = font, BOR
    if fmt:
        x.number_format = fmt
    if fill:
        x.fill = fill
    if wrap or center:
        x.alignment = Alignment(wrap_text=wrap, vertical='center', horizontal='center' if center else None)
    return x


def anadir(wb, bc, pos=1):
    ws = wb.create_sheet('PDF vs BC', pos)
    ws['A1'] = 'Liquidación ECI Portugal, agosto 2026: columnas del PDF de ECI frente a lo que calcula BC'
    ws['A1'].font = T
    ws['A2'] = ('Valores del PDF en azul, copiados de los tres PDF de Portugal. Valores de BC por fórmula: suma de las filas de Portugal del informe de BC '
                '(centros 0140 + 0142 + 0143), que están abajo en «Datos de BC por centro».')
    ws['A2'].font = NOTA
    r0 = 4
    grupos = [('601', '601 mujer'), ('602', '602 hombre'), ('696', '696 niños'), (None, 'Total Portugal')]
    for j, h in enumerate(['Columna del PDF de ECI', 'Nombre en el PDF', 'Campo equivalente en BC'], 1):
        ws.merge_cells(start_row=r0, start_column=j, end_row=r0 + 1, end_column=j)
        _c(ws, r0, j, h, CABF, fill=CAB, wrap=True)
        _c(ws, r0 + 1, j, None)
    for g, (_, nombre) in enumerate(grupos):
        c0 = 4 + g * 3
        ws.merge_cells(start_row=r0, start_column=c0, end_row=r0, end_column=c0 + 2)
        _c(ws, r0, c0, nombre, CABF, fill=CAB, center=True)
        for k, sh in enumerate(['PDF ECI', 'BC', 'Diferencia BC − PDF']):
            _c(ws, r0 + 1, c0 + k, sh, B, fill=SUB, wrap=True, center=True)
    cc = 4 + len(grupos) * 3
    ws.merge_cells(start_row=r0, start_column=cc, end_row=r0 + 1, end_column=cc)
    _c(ws, r0, cc, 'Comentario', CABF, fill=CAB, wrap=True)
    _c(ws, r0 + 1, cc, None)
    for j, w in enumerate([26, 18, 26] + [12, 12, 12] * 4 + [70], 1):
        ws.column_dimensions[L(j)].width = w
    ws.row_dimensions[r0 + 1].height = 30

    # bloque de datos de BC (base de las fórmulas)
    rd0 = r0 + 2 + len(FILAS) + 5
    ws.cell(row=rd0 - 1, column=1, value='Datos de BC por centro (filas de Portugal de «liquidación mensual ECI mes 8 datos.xlsx», 02/10/2026)').font = B
    for j, h in enumerate(['Centro', 'UNECO', 'Estado'] + BC_COLS, 1):
        _c(ws, rd0, j, h, CABF, fill=CAB, wrap=True)
    ws.row_dimensions[rd0].height = 30
    for i, (_, b) in enumerate(bc.iterrows(), rd0 + 1):
        vals = [str(b['Código centro']), str(b['Código UNECO candidato']), b['Estado comparación operativa']] + [float(b[c]) for c in BC_COLS]
        for j, v in enumerate(vals, 1):
            _c(ws, i, j, v, fmt=EUR if j > 3 else None)
    rd1, rd2 = rd0 + 1, rd0 + len(bc)
    colbc = {c: L(4 + k) for k, c in enumerate(BC_COLS)}

    def sbc(campo, u):
        col = colbc[campo]
        return f'SUMIFS(${col}${rd1}:${col}${rd2},$B${rd1}:$B${rd2},"{u}")'

    fila = {key: r0 + 2 + k for k, (_, _, _, key, _, _) in enumerate(FILAS)}
    for i, (concepto, nombre_pdf, campo_bc, key, campo, comentario) in enumerate(FILAS, r0 + 2):
        _c(ws, i, 1, concepto, B)
        _c(ws, i, 2, nombre_pdf)
        _c(ws, i, 3, campo_bc)
        fmt, fdif = (PCT, PCT) if key == 'pct' else (EUR, DIF)
        for g, (u, _) in enumerate(grupos):
            c0 = 4 + g * 3
            font = N if u else B
            if u:
                _c(ws, i, c0, PDF[u][key], AZ, fmt)
                if campo is None:
                    v = 'no existe'
                elif campo == 'PCT':
                    v = f'=IFERROR({sbc("Comisión esperada", u)}/{sbc("Base imponible comisión", u)},"")'
                else:
                    v = f'={sbc(campo, u)}'
            else:  # total
                if key == 'pct':
                    _c(ws, i, c0, f'={L(c0)}{fila["part"]}/{L(c0)}{fila["neta"]}', B, fmt)
                    v = f'=IFERROR({L(c0 + 1)}{fila["part"]}/{L(c0 + 1)}{fila["neta"]},"")'
                else:
                    _c(ws, i, c0, '=' + '+'.join(f'{L(4 + k * 3)}{i}' for k in range(3)), B, fmt)
                    v = 'no existe' if campo is None else '=' + '+'.join(f'N({L(5 + k * 3)}{i})' for k in range(3))
            _c(ws, i, c0 + 1, v, font, fmt)
            _c(ws, i, c0 + 2, f'=IF(ISNUMBER({L(c0 + 1)}{i}),{L(c0 + 1)}{i}-{L(c0)}{i},"")', font, fdif)
        x = _c(ws, i, cc, comentario, NOTA)
        x.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[i].height = 30
    rf = r0 + 2 + len(FILAS) + 1
    for k, t in enumerate([
        'Lectura: la venta con IVA coincide al céntimo. A partir de ahí BC se separa porque calcula la comisión y el neto sobre el importe con IVA, '
        'mientras ECI quita antes el IVA portugués del 23 % y Toni Pons factura sin IVA.',
        'Por eso el «Importe neto esperado» de BC queda 807,22 € por encima del total de las facturas de Portugal (260007291, 260007292 y 260007293 = 3.466,96 €).',
    ]):
        ws.cell(row=rf + k, column=1, value=t).font = NOTA
    ws.freeze_panes = 'D6'
    ws.sheet_view.showGridLines = False
    return ws
