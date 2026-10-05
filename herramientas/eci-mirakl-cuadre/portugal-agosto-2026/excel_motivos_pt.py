# -*- coding: utf-8 -*-
"""Excel de detalle del informe de motivos de Portugal (agosto 2026)."""
import json
import sys

import openpyxl
import pandas as pd
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')
SC, OUT, ECI = sys.argv[1], sys.argv[2], sys.argv[3]
rec = pd.read_parquet(SC + '/pt/reconstruccion_pt.parquet')
bc = pd.DataFrame(json.load(open(SC + '/pt/bc_informe_pt.json', encoding='utf-8')))

# ---------------------------------------------------------------- ajustes con el detalle de BC (05/10/2026)
# Detalle de la comparación operativa exportado de BC para las filas con aviso o diferencia con Mirakl.
DETALLES = {('0142', '601'): 'detalle gaia.xlsx', ('0143', '601'): 'detalle gaia 143.xlsx', ('0143', '602'): '143 602.xlsx'}
mk = pd.read_parquet(SC + '/pt/pt_mirakl_export.parquet').set_index('N.º de asiento de pedido')
lg = pd.read_parquet(SC + '/pt/pt_log_conector_bc_0109.parquet').set_index('Mirakl Line Id')


def _num(v):
    v = pd.to_numeric(v, errors='coerce')
    return 0.0 if pd.isna(v) else float(v)


def base(a):
    m = mk.loc[a]
    d8 = lambda c: pd.to_datetime(m[c] or None, format='%Y%m%d', errors='coerce')
    dt = lambda c: pd.to_datetime(m[c] or None, format='%d/%m/%Y %H:%M:%S', errors='coerce')
    return dict(asiento=a, pedido=m['Número de pedido'], uneco=m['UNECO_inferido_por_descripcion'], ean=m['SKU del producto'],
                descripcion=str(m['Detalles']).split(' (SKU')[0][:70], estado=m['Estado'], cantidad=_num(m['Cantidad']),
                importe=_num(m['Importe']), reembolso=_num(m['Importe total reembolsado (impuestos incluidos)']),
                creado=dt('Fecha de creación'), debito=dt('Fecha de débito al cliente'), cumplimentacion=d8('Fecha de cumplimentación'),
                talon_cumpl=m['Talón de cumplimentación'], abono=d8('Fecha de abono 1'), talon_abono=m['Talón de abono 1'],
                recogida=d8('Fecha de recogida 1'), conector=lg['Processing Status'].get(a, ''), pedido_bc=lg['BC Sales Order No.'].get(a, ''))


GAIA = ('Devolución en la tienda de Gaia el 10/08 según el EDI. En Mirakl el pedido sigue «Recibido»: sin devolución, incidencia, talón de abono '
        'ni mensajes. BC la vincula por el talón de venta y no puede procesarla. Pregunta a ECI.')
extra = [
    ('00401430751578320260801170207_1-A-1', 'Devolución', '0142', -59.95, GAIA),
    ('00401430751578320260801170207_2-A-1', 'Devolución', '0142', -69.95, GAIA),
    ('00401430751578320260801170207_2-A-2', 'Devolución', '0142', -89.95, GAIA),
    ('00401430764858020260717201700_1-A-1', 'Devolución', '0143', -79.95,
     'El EDI del 05/08 abona las dos líneas del pedido; en Mirakl solo se reembolsó la 2-A-1 (75 €) y esta sigue «Recibido». '
     'BC la vincula por el talón de venta y no puede procesarla. Pregunta a ECI.'),
    ('00401430757641420260826172900_1-A-1', 'Venta finalizada', '0143', 65.00,
     'Corte de fin de mes: recepción el 31/08 a las 23:52. Mirakl fecha el talón de cumplimentación el 01/09, pero ECI la liquida en agosto (EDI del 31/08).'),
]
rec = pd.concat([rec, pd.DataFrame([{**base(a), 'concepto': c, 'centro': ce, 'importe_edi': imp, 'vinculo': 'Vinculada',
                                     'motivo': mot + ' (identificada con el detalle de BC)'} for a, c, ce, imp, mot in extra])], ignore_index=True)
cuarta = (rec['asiento'] == '00401430752499520260615175342_1-A-1') & (rec['concepto'] == 'Venta finalizada')
rec.loc[cuarta, 'vinculo'] = 'Sin vincular'
rec.loc[cuarta, 'motivo'] = ('Pedido del 15/06 que BC tiene como venta 124836 desde junio. Lo más probable: el conector dejó de actualizarlo (pedido anterior al 29/06) '
                             'y no tiene el talón de cumplimentación del 25/08 con el que BC vincula las ventas. Por confirmar con IT.')

# procesado y documento BC de cada línea, según el detalle de BC
rec['procesado_bc'], rec['doc_bc'] = '', ''
for (ce, u), f in DETALLES.items():
    det = pd.read_excel(ECI + '/' + f)
    for _, d in det.iterrows():
        tipo = d['Tipo movimiento']
        sel = (rec['centro'] == ce) & (rec['uneco'] == u) & (rec['concepto'] == tipo)
        if isinstance(d['Id. línea Mirakl'], str):
            sel &= rec['asiento'] == d['Id. línea Mirakl']
        else:
            sel &= (rec['vinculo'] == 'Sin vincular') & (rec['ean'].astype(str) == str(int(d['EAN']))) & ((rec['importe_edi'] - d['Importe liquidable']).abs() < 0.005)
        assert sel.sum() == 1, (f, d['EAN'], d['Importe liquidable'], int(sel.sum()))
        rec.loc[sel, 'procesado_bc'] = 'Sí' if d['Procesado en BC'] == 1 else 'No'
        rec.loc[sel, 'doc_bc'] = '' if pd.isna(d['N.º documento BC']) else str(int(d['N.º documento BC']))
    n_bc = int((det['Tipo movimiento'].isin(['Venta finalizada', 'Devolución'])).sum())
    n_rec = int(((rec['centro'] == ce) & (rec['uneco'] == u) & rec['concepto'].isin(['Venta finalizada', 'Devolución'])).sum())
    assert n_bc == n_rec, (f, n_bc, n_rec)

F = 'Arial'
CAB = PatternFill('solid', fgColor='1F4E78')
CABF = Font(name=F, size=10, bold=True, color='FFFFFF')
N = Font(name=F, size=10)
B = Font(name=F, size=10, bold=True)
T = Font(name=F, size=14, bold=True, color='1F4E78')
NOTA = Font(name=F, size=9, italic=True, color='595959')
th = Side(style='thin', color='BFBFBF')
BOR = Border(left=th, right=th, top=th, bottom=th)
EUR = '#,##0.00 "€";-#,##0.00 "€";"-"'
DIF = '#,##0.00 "€";[Red]-#,##0.00 "€";"-"'
DIA = 'dd/mm/yyyy'
FH = 'dd/mm/yyyy hh:mm'


def cab(ws, r, heads, widths=None):
    for j, h in enumerate(heads, 1):
        c = ws.cell(row=r, column=j, value=h)
        c.fill, c.font, c.border = CAB, CABF, BOR
        c.alignment = Alignment(wrap_text=True, vertical='center')
        if widths:
            ws.column_dimensions[get_column_letter(j)].width = widths[j - 1]
    ws.row_dimensions[r].height = 30


def cel(ws, r, c, v, fmt=None, font=N, wrap=False):
    if isinstance(v, float) and pd.isna(v):
        v = None
    if isinstance(v, pd.Timestamp):
        v = None if pd.isna(v) else v.to_pydatetime()
    x = ws.cell(row=r, column=c, value=v)
    x.font, x.border = font, BOR
    if fmt:
        x.number_format = fmt
    if wrap:
        x.alignment = Alignment(wrap_text=True, vertical='top')
    return x


wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Resumen por fila BC'
ws2 = wb.create_sheet('Líneas EDI reconstruidas')
ws3 = wb.create_sheet('Identificadas con detalle BC')
ws6 = wb.create_sheet('Devoluciones en tienda')
ws4 = wb.create_sheet('Comisión e IVA')
ws5 = wb.create_sheet('Preguntas para IT y ECI')

# ---------------------------------------------------------------- hoja 2: líneas reconstruidas
cols = ['asiento', 'pedido', 'centro', 'uneco', 'concepto', 'vinculo', 'importe_edi', 'cantidad', 'importe', 'reembolso', 'estado',
        'ean', 'descripcion', 'creado', 'debito', 'cumplimentacion', 'talon_cumpl', 'abono', 'talon_abono', 'recogida', 'conector', 'pedido_bc',
        'procesado_bc', 'doc_bc', 'motivo']
heads = ['N.º asiento Mirakl', 'Nº pedido Mirakl', 'Centro EDI', 'UNECO', 'Concepto EDI', 'Vínculo con Mirakl', 'Importe EDI', 'Uds EDI',
         'Importe línea Mirakl', 'Reembolso Mirakl', 'Estado Mirakl (21/09)', 'EAN', 'Descripción', 'Creado', 'Cobrado', 'Cumplimentación',
         'Talón cumplim.', 'Abono', 'Talón abono', 'Recogida', 'Conector (01/09)', 'Pedido BC', 'Procesado en BC (detalle 05/10)', 'Documento BC', 'Motivo']
orden = {'Venta finalizada': 1, 'Excluida del finalizado': 2, 'Devolución': 3, 'No es devolución': 4, 'Rescisión': 5, 'En proceso': 6, 'Fuera del EDI de agosto': 7}
rec = rec.assign(_o=rec['concepto'].map(orden)).sort_values(['_o', 'centro', 'uneco', 'vinculo', 'asiento']).drop(columns='_o')
ws2['A1'] = 'Líneas EDI de agosto de 2026 de Portugal reconstruidas desde el export de Mirakl (21/09) con las reglas verificadas'
ws2['A1'].font = T
ws2['A2'] = ('Una fila por línea EDI. «Importe EDI» lleva el signo de la liquidación (devoluciones y rescisiones en negativo). '
             'En proceso y rescisión no entran en la liquidación. Las 4 devoluciones sin procesar y la venta de 65 € de hombre no salen de las reglas: '
             'se han añadido con el detalle de la comparación operativa de BC (05/10), igual que «Procesado en BC» y «Documento BC» de las filas 0142/601, 0143/601 y 0143/602.')
ws2['A2'].font = NOTA
cab(ws2, 4, heads, [36, 34, 8, 7, 18, 12, 11, 6, 11, 10, 15, 15, 34, 15, 15, 12, 11, 12, 11, 12, 18, 9, 12, 10, 80])
fmt = {'importe_edi': DIF, 'importe': EUR, 'reembolso': EUR, 'creado': FH, 'debito': FH, 'cumplimentacion': DIA, 'abono': DIA, 'recogida': DIA}
for i, rr in enumerate(rec[cols].itertuples(index=False), 5):
    for j, (k, v) in enumerate(zip(cols, rr), 1):
        cel(ws2, i, j, v, fmt.get(k))
n2 = 4 + len(rec)
ws2.freeze_panes = 'B5'
ws2.auto_filter.ref = f'A4:{get_column_letter(len(cols))}{n2}'
ws2.conditional_formatting.add(f'A5:{get_column_letter(len(cols))}{n2}',
                               FormulaRule(formula=[f'${get_column_letter(cols.index("procesado_bc") + 1)}5="No"'], fill=PatternFill('solid', fgColor='FFEB9C', bgColor='FFEB9C')))
for txt, color in [('Sin vincular', 'FFC7CE'), ('Vinculada', 'C6EFCE')]:
    ws2.conditional_formatting.add(f'A5:{get_column_letter(len(cols))}{n2}',
                                   FormulaRule(formula=[f'$F5="{txt}"'], fill=PatternFill('solid', fgColor=color, bgColor=color)))
for txt in ['Excluida del finalizado', 'No es devolución', 'Fuera del EDI de agosto']:
    ws2.conditional_formatting.add(f'A5:{get_column_letter(len(cols))}{n2}',
                                   FormulaRule(formula=[f'$E5="{txt}"'], fill=PatternFill('solid', fgColor='E7E6E6', bgColor='E7E6E6')))
R = lambda col: f"'Líneas EDI reconstruidas'!${get_column_letter(cols.index(col) + 1)}$5:${get_column_letter(cols.index(col) + 1)}${n2}"

# ---------------------------------------------------------------- hoja 1: resumen por fila BC
ws['A1'] = 'Liquidación ECI Portugal, agosto de 2026: informe de BC frente a la reconstrucción desde Mirakl'
ws['A1'].font = T
ws['A2'] = ('Fuentes: «liquidación mensual ECI mes 8 datos.xlsx» (BC, 02/10/2026), «pedidos 2026.xlsx» (Mirakl, 21/09/2026), log del conector (01/09/2026), '
            'PDF de ECI de Portugal (8 - AGOST) y detalle de la comparación operativa de BC (05/10/2026). Columnas de reconstrucción con fórmulas sobre la hoja de líneas.')
ws['A2'].font = NOTA
h = ['Centro', 'UNECO', 'Estado BC', 'Liquidable BC', 'Finalizado BC', 'Finalizado reconstruido', 'Dif. finalizado', 'Devoluciones BC',
     'Devoluciones reconstruidas', 'Dif. devoluciones', 'Líneas sin vincular BC', 'Sin vincular reconstruidas', 'Tramos sin procesar BC',
     'Sin procesar (detalle BC)', 'Motivo']
cab(ws, 4, h, [8, 7, 11, 13, 13, 14, 12, 13, 14, 12, 10, 11, 10, 11, 95])
motivos = {
    ('0140', '601'): 'Cuadrado. 9 devoluciones hechas en la tienda de Lisboa (cajas 0247, 0558-0561), todas con tramo procesado. ECI las liquida dentro de 0143.',
    ('0140', '696'): 'Cuadrado. 1 devolución de niños en Lisboa (talón 03594705) en Incidencia abierta: ECI la liquida por el importe completo aunque Mirakl no haya reembolsado.',
    ('0142', '601'): ('Aviso. 9 devoluciones de Gaia procesadas (518,12), 7 de ellas con la incidencia aún abierta en Mirakl. 1 sin vincular de 34,95: pedido de junio que el conector dejó de actualizar. '
                      '3 sin procesar (219,85): pedido 00401430751578320260801170207, devuelto en Gaia el 10/08 según el EDI y sin devolución en Mirakl.'),
    ('0143', '601'): ('Aviso. 6 líneas sin vincular (neto +68,91): 4 ventas de pedidos de enero a junio (220,82) y 2 devoluciones de pedidos de junio congelados en el conector (151,91). '
                      '1 devolución sin procesar (79,95): línea 00401430764858020260717201700_1-A-1, abonada por ECI y sin reembolso en Mirakl. 1 producto sin relacionar sin localizar.'),
    ('0143', '602'): 'Cuadrado. Incluye la venta de 65,00 del pedido 00401430757641420260826172900: recepción el 31/08 a las 23:52, que Mirakl fecha el 01/09 y ECI liquida en agosto.',
    ('0143', '696'): 'Cuadrado. Sin incidencias.',
}
r = 5
for _, b in bc.iterrows():
    c, u = str(b['Código centro']), str(b['Código UNECO candidato'])
    cel(ws, r, 1, c)
    cel(ws, r, 2, u)
    cel(ws, r, 3, b['Estado comparación operativa'])
    cel(ws, r, 4, float(b['Importe liquidable']), EUR)
    cel(ws, r, 5, float(b['Importe finalizado']), EUR)
    cel(ws, r, 6, f'=SUMIFS({R("importe_edi")},{R("centro")},$A{r},{R("uneco")},$B{r},{R("concepto")},"Venta finalizada")', EUR)
    cel(ws, r, 7, f'=ROUND(E{r}-F{r},2)', DIF)
    cel(ws, r, 8, float(b['Importe devoluciones']), EUR)
    cel(ws, r, 9, f'=SUMIFS({R("importe_edi")},{R("centro")},$A{r},{R("uneco")},$B{r},{R("concepto")},"Devolución")', EUR)
    cel(ws, r, 10, f'=ROUND(H{r}-I{r},2)', DIF)
    cel(ws, r, 11, int(b['Líneas EDI liquidables sin vincular']))
    cel(ws, r, 12, f'=COUNTIFS({R("centro")},$A{r},{R("uneco")},$B{r},{R("vinculo")},"Sin vincular")')
    cel(ws, r, 13, int(b['Tramos de origen sin procesar en BC']))
    cel(ws, r, 14, f'=COUNTIFS({R("centro")},$A{r},{R("uneco")},$B{r},{R("procesado_bc")},"No",{R("vinculo")},"Vinculada")')
    cel(ws, r, 15, motivos.get((c, u), ''), wrap=True)
    ws.row_dimensions[r].height = 45
    r += 1
cel(ws, r, 1, 'Total', font=B)
for col in 'DEFGHIJ':
    cel(ws, r, ord(col) - 64, f'=SUM({col}5:{col}{r - 1})', EUR if col not in 'GJ' else DIF, B)
for col in 'KLMN':
    cel(ws, r, ord(col) - 64, f'=SUM({col}5:{col}{r - 1})', '0', B)
tot = r
ws.conditional_formatting.add(f'A5:O{r - 1}', FormulaRule(formula=['$C5="Aviso"'], fill=PatternFill('solid', fgColor='FFEB9C', bgColor='FFEB9C')))
ws.conditional_formatting.add(f'A5:O{r - 1}', FormulaRule(formula=['$C5="Cuadrado"'], fill=PatternFill('solid', fgColor='C6EFCE', bgColor='C6EFCE')))
r += 2
ws.cell(row=r, column=1, value='Cuadre con los PDF de ECI (ECI solo liquida el centro 0143 y mete dentro las devoluciones de las tiendas 0140 y 0142)').font = B
r += 1
cab(ws, r, ['UNECO', '', '', 'Liquidable BC (0140+0142+0143)', 'Total venta PDF ECI', 'Diferencia'])
r += 1
pdf = {'601': 4922.72, '602': 981.56, '696': 101.87}
r0 = r
for u, v in pdf.items():
    cel(ws, r, 1, u)
    cel(ws, r, 4, f'=SUMIFS($D$5:$D${tot - 1},$B$5:$B${tot - 1},$A{r})', EUR)
    cel(ws, r, 5, v, EUR)
    cel(ws, r, 6, f'=ROUND(D{r}-E{r},2)', DIF)
    r += 1
cel(ws, r, 1, 'Total', font=B)
cel(ws, r, 4, f'=SUM(D{r0}:D{r - 1})', EUR, B)
cel(ws, r, 5, f'=SUM(E{r0}:E{r - 1})', EUR, B)
cel(ws, r, 6, f'=ROUND(D{r}-E{r},2)', DIF, B)
ws.freeze_panes = 'C5'
ws.sheet_view.showGridLines = False

# ---------------------------------------------------------------- hoja 3: pendiente de identificar
ws3['A1'] = 'Líneas que no salían de las reglas, identificadas con el detalle de la comparación operativa de BC (05/10/2026)'
ws3['A1'].font = T
cab(ws3, 3, ['Fila BC', 'Línea de Mirakl', 'Importe EDI', 'Qué es', 'Estado', 'Qué hacer'], [10, 40, 11, 80, 14, 60])
pend = [
    ('0142/601', '00401430751578320260801170207_1-A-1, _2-A-1 y _2-A-2', -219.85,
     'Tres pares de alpargatas pedidos el 01/08 y entregados en casa el 07/08. El EDI (informe 20260811035705067032352) los da por devueltos en la tienda de Gaia el 10/08. '
     'En Mirakl el pedido sigue «Recibido», sin devolución, incidencia, talón de abono ni mensajes. BC los vincula por el talón de venta 07515783 y no puede procesarlos.',
     'Tramo sin procesar', 'ECI: confirmar si la devolución se hizo y pasar el talón de abono. Si se hizo, registrarla en Mirakl; si no, regularizar 219,85 €.'),
    ('0143/601', '00401430764858020260717201700_1-A-1', -79.95,
     'Menorquinas de mujer, pedido del 17/07. El EDI del 05/08 (informe 20260806042406067032352) abona las dos líneas del pedido. En Mirakl solo está reembolsada la 2-A-1 (75 €, talón 92363569, '
     'documento BC 6245); la 1-A-1 sigue «Recibido».',
     'Tramo sin procesar', 'ECI: confirmar si se devolvió también la 1-A-1. Si es así, registrarla en Mirakl; si no, regularizar 79,95 €.'),
    ('0143/601', '00401430752499520260615175342_1-A-1', 59.95,
     'Cuarta venta sin vincular. Pedido del 15/06, entregado el 19/06 y cumplimentado por ECI el 25/08 (talón 94593704). BC la tiene como venta 124836 desde junio. '
     'Lo más probable es que el conector no tenga su talón de cumplimentación, porque dejó de actualizar los pedidos creados hasta el 28/06, y BC vincula las ventas por ese talón.',
     'Sin vincular', 'IT: confirmar el motivo y vincularla.'),
    ('0143/602', '00401430757641420260826172900_1-A-1', 65.00,
     'Zuecos de hombre entregados el 31/08 a las 18:47, con la recepción confirmada a las 23:52. Mirakl fecha el talón de cumplimentación el 01/09, pero ECI la liquida en agosto '
     '(EDI del 31/08). BC la tiene procesada (documento 145645).',
     'Correcto', 'Nada. Tenerlo en cuenta al cuadrar septiembre: Mirakl la pondrá en septiembre aunque ECI ya la haya liquidado.'),
    ('0143/601', 'Producto sin relacionar (1)', None,
     'No está entre las líneas liquidables del detalle: todas tienen producto de BC. Probablemente es una línea no liquidable (en proceso o rescisión); en España aparece en una fila con una sola rescisión.',
     'Sin localizar', 'IT: indicar qué línea o EAN es y si impide vincular.'),
]
for i, p in enumerate(pend, 4):
    for j, v in enumerate(p, 1):
        cel(ws3, i, j, v, DIF if j == 3 else None, wrap=True)
    ws3.row_dimensions[i].height = 90

# ---------------------------------------------------------------- hoja 6: devoluciones en tienda
ws6['A1'] = 'Devoluciones en tienda de Gaia: mensajes R01, R02 y R03 de Mirakl'
ws6['A1'].font = T
ws6['A2'] = ('R01: el cliente devuelve en la tienda y le abonan por el TPV (ECI la liquida en ese mes). R02: la mercancía llega al almacén de MRW. '
             'R03: Toni Pons la recoge y Mirakl crea el reembolso. Fechas de los mensajes de Mirakl consultados el 05/10/2026.')
ws6['A2'].font = NOTA
cab(ws6, 4, ['Línea de Mirakl', 'Talón de abono', 'Importe', 'R01', 'R02', 'R03', 'Reembolso Mirakl', 'En el EDI de agosto', 'Procesado en BC', 'Documento BC'],
    [38, 13, 10, 11, 11, 11, 16, 14, 12, 12])
tienda = [
    ('00401430770000620260726233949_1-A-1', '02695702', 59.95, '06/08/2026', '29/09/2026', '', 'No (incidencia abierta)', 'Sí', 'Sí', '6265'),
    ('00401430766772720260803145554_1-A-1', '02654478', 79.95, '08/08/2026', '24/09/2026', '', 'No (incidencia abierta)', 'Sí', 'Sí', '6255'),
    ('00401430759576120260802110116_1-A-1', '06171499', 29.95, '13/08/2026', '', '', 'No (incidencia abierta)', 'Sí', 'Sí', '6234'),
    ('00401430759576120260802110116_1-A-2', '06171499', 29.95, '13/08/2026', '', '', 'No (incidencia abierta)', 'Sí', 'Sí', '6234'),
    ('00401430759559920260731144944_1-A-1', '02655550', 79.95, '17/08/2026', '', '', 'No (incidencia abierta)', 'Sí', 'Sí', '6233'),
    ('00401430768572820260811204517_1-A-1', '02655611', 69.95, '17/08/2026', '23/09/2026', '', 'No (incidencia abierta)', 'Sí', 'Sí', '6263'),
    ('00401430766919320260821222205_1-A-1', '02667327', 27.96, '28/08/2026', '', '', 'No (incidencia abierta)', 'Sí', 'Sí', '6256'),
    ('00401430763871220260711145749_1-A-1', '06166877', 79.95, '16/07/2026', '23/07/2026', '24/07/2026', 'Sí, 24/07 (hora de la R03)', 'No (R01 de julio)', '', ''),
    ('00401430750813420260714121528_1-A-1', '02652811', 59.95, '27/07/2026', '24/08/2026', '31/08/2026', 'Sí, 31/08 (hora de la R03)', 'No (R01 de julio)', '', ''),
]
for i, t in enumerate(tienda, 5):
    for j, v in enumerate(t, 1):
        cel(ws6, i, j, v, EUR if j == 3 else None)
r = 5 + len(tienda) + 1
for t in ['Las 7 líneas de agosto no tienen R03 (ni reembolso en Mirakl) y aun así están en el EDI de agosto y procesadas en BC: ECI liquida con la R01 y BC no espera al reembolso.',
          'Las dos de julio tienen la R03 y el reembolso en julio o agosto, pero no están en la liquidación de agosto: ECI las liquida en el mes de la R01.',
          'En agosto la mercancía tardó entre 5 y 7 semanas en llegar al almacén (R02), y en tres casos todavía no ha llegado.']:
    ws6.cell(row=r, column=1, value=t).font = NOTA
    r += 1

# ---------------------------------------------------------------- hoja 4: comisión e IVA
ws4['A1'] = 'Comisión de ECI y neto a facturar: cálculo de ECI (PDF) frente a cálculo de BC'
ws4['A1'].font = T
ws4['A2'] = ('ECI: neta = total / 1,23 (IVA portugués); participación = 29 % de la neta; Toni Pons factura la neta menos la participación SIN IVA. '
             'BC: comisión = 29 % de la base con IVA; neto esperado = liquidable − comisión.')
ws4['A2'].font = NOTA
cab(ws4, 4, ['UNECO', 'Total venta ECI', 'Neta ECI', 'IVA 23 %', 'Participación ECI 29 %', 'A facturar ECI (sin IVA)', 'Liquidable BC', 'Base comisión BC',
             'Comisión esperada BC', 'Neto esperado BC', 'Dif. neto BC − a facturar', 'Dif. comisión BC − ECI'], [8, 13, 12, 11, 13, 14, 13, 13, 13, 13, 14, 13])
ecis = {'601': (4922.72, 4002.21, 920.51, 1160.64, 2841.57), '602': (981.56, 798.02, 183.54, 231.43, 566.59), '696': (101.87, 82.82, 19.05, 24.02, 58.80)}
bcu = bc.assign(u=bc['Código UNECO candidato'].astype(str)).groupby('u')[['Importe liquidable', 'Base imponible comisión', 'Comisión esperada', 'Importe neto esperado']].sum()
r = 5
for u, (t_, n_, iv, pa, af) in ecis.items():
    vals = [u, t_, n_, iv, pa, af] + [round(float(bcu.loc[u, c]), 2) for c in ['Importe liquidable', 'Base imponible comisión', 'Comisión esperada', 'Importe neto esperado']]
    for j, v in enumerate(vals, 1):
        cel(ws4, r, j, v, None if j == 1 else EUR)
    cel(ws4, r, 11, f'=ROUND(J{r}-F{r},2)', DIF)
    cel(ws4, r, 12, f'=ROUND(I{r}-E{r},2)', DIF)
    r += 1
cel(ws4, r, 1, 'Total', font=B)
for col in 'BCDEFGHIJKL':
    cel(ws4, r, ord(col) - 64, f'=SUM({col}5:{col}{r - 1})', EUR if col not in 'KL' else DIF, B)
r += 2
for t in ['Lectura: el neto esperado de BC supera en el total lo que realmente se factura a ECI. La diferencia = IVA portugués que BC no descuenta − exceso de comisión que BC calcula sobre el importe con IVA.',
          'Además, BC deja fuera de la base de comisión las líneas sin vincular pero las suma enteras en el neto esperado (601: 33,96 de neto sin comisión).',
          '«Comisión real Mirakl» vale 0 en todas las filas: ECI no carga su participación en Mirakl (las 27.730 líneas del export tienen comisión 0), la cobra en la liquidación mensual.']:
    ws4.cell(row=r, column=1, value=t).font = NOTA
    r += 1

# ---------------------------------------------------------------- hoja 5: preguntas
ws5['A1'] = 'Preguntas y peticiones'
ws5['A1'].font = T
cab(ws5, 3, ['Para', 'Tema', 'Pregunta o petición'], [8, 30, 120])
qs = [
    ('IT', 'Comisión e IVA de Portugal', 'Calcular la comisión esperada sobre la venta neta sin IVA del país (23 % en Portugal) y el neto esperado sin IVA en Portugal (Toni Pons factura a ECI Portugal sin IVA). Hoy el neto esperado de Portugal sale 807,22 € por encima de lo facturable.'),
    ('IT', 'Comisión de niños en España', 'BC aplica el 29 % a las filas 696 de España y ECI aplica el 30 % en 0090/696. Revisar la tabla de porcentajes por país y UNECO.'),
    ('IT', 'Pedidos congelados en el conector', 'Los pedidos creados hasta el 28/06 dejaron de actualizarse en el conector: sus devoluciones posteriores no generan tramo y salen como líneas EDI sin vincular. Forzar una relectura de esos pedidos o crear el abono a mano. En julio hay al menos 12 devoluciones más en esta situación.'),
    ('IT', 'Ventas de pedidos anteriores al conector', 'ECI sigue liquidando entregas de pedidos de enero a mayo (registrados en SAP). Definir cómo se registran en BC: vínculo manual o documento manual.'),
    ('IT', 'Venta 124836 sin vincular', 'La venta del pedido 00401430752499520260615175342 (15/06) está en BC desde junio, pero su línea EDI del 25/08 no se vincula. Confirmar si es porque el conector no tiene su talón de cumplimentación (pedido congelado) y vincularla.'),
    ('IT', 'Devoluciones antes del reembolso', 'BC crea la devolución con la incidencia abierta, antes de que Mirakl reembolse (el reembolso llega con la R03, semanas después). Revisar cómo se corrige si el importe reembolsado es distinto (p. ej. 00401430763049820260801022856: 2 líneas de 150 € con 75 € reembolsados en cada una).'),
    ('IT', 'Producto sin relacionar', 'Indicar qué línea o EAN cuenta como «producto sin relacionar» en 0143/601 y si bloquea la vinculación.'),
    ('IT', 'Cantidades devueltas en España', 'Algunas filas de España tienen cantidades devueltas imposibles (0037/601: -223 uds en 28 líneas; 0003/601: -214 en 16). Revisar el campo cantidad en la importación EDI.'),
    ('ECI', 'Liquidación por centro', 'Los PDF dicen «facturar independientemente por centro» pero solo traen 0143, con las devoluciones de las tiendas de Lisboa y Gaia dentro. Confirmar que se factura por departamento consolidado.'),
    ('ECI', 'Devolución en Gaia sin registrar en Mirakl', 'El EDI del 10/08 abona en la tienda de Gaia el pedido 00401430751578320260801170207 (3 líneas, 219,85 €), que en Mirakl sigue «Recibido» sin devolución ni mensajes. Confirmar si se hizo y pasar el talón de abono; si se hizo, registrarla en Mirakl; si no, regularizar.'),
    ('ECI', 'Abono online sin reembolso en Mirakl', 'El EDI del 05/08 abona la línea 00401430764858020260717201700_1-A-1 (79,95 €) junto a su hermana de 75 €, pero en Mirakl solo se reembolsó la de 75 €. Confirmar si se devolvió.'),
    ('ECI', 'Plazo de las devoluciones en tienda', 'La mercancía devuelta en tienda tarda de 5 a 7 semanas en llegar al almacén (R02). Mientras tanto Mirakl no refleja el reembolso.'),
]
for i, q in enumerate(qs, 4):
    for j, v in enumerate(q, 1):
        cel(ws5, i, j, v, wrap=True)
    ws5.row_dimensions[i].height = 45

sys.path.insert(0, SC + "/pt")
from hoja_pdf_vs_bc import anadir
anadir(wb, bc, 1)

for w in (ws, ws3, ws4, ws5, ws6):
    w.sheet_view.showGridLines = False
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('guardado', OUT, '| líneas', len(rec))
