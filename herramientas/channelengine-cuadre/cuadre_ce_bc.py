# -*- coding: utf-8 -*-
"""
Cuadre de pedidos y devoluciones ChannelEngine (CE) vs Business Central (BC).

Entradas
  --bc   Excel con el log de la integración BC (hoja 'Cruce pedidos', cabecera
         con 'Channel Order No', 'Estado CE', 'Total BC', 'Estado proceso BC'…)
  --ce   CSV 'orders-report-*.csv' descargado del backoffice de ChannelEngine
         (separador ';', una fila por línea de pedido, primera línea 'sep=;')
Salida
  --salida  Excel de cuadre con las hojas Resumen, Cruce pedidos, Devoluciones,
            Pedidos CE, Líneas CE y Log BC. Las columnas de cruce son fórmulas,
            y los colores son formato condicional sobre la columna Resultado.

Uso
  python cuadre_ce_bc.py --bc "Cruce_pedidos_ChannelEngine_agosto_2026.xlsx" ^
      --ce "orders-report-xxx.csv" --periodo 2026-08 --salida "Cuadre_CE_vs_BC_agosto_2026.xlsx"
"""
import argparse
import calendar
import shutil
import zipfile
import datetime as dt
import sys
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FUENTE = "Arial"
MAXF = 2000  # filas máximas que cubren las fórmulas (rangos acotados, no columnas enteras)

# Paleta (relleno de fila por Resultado)
COLOR = {
    "Coincide": "C6EFCE",
    "Coincide con diferencia de importe": "FFEB9C",
    "Devolución con pedido BC": "D9C3EA",
    "Devolución omitida en BC (sin pedido)": "DDEBF7",
    "Omitido en BC (sin pedido)": "DDEBF7",
    "Error BC (sin pedido BC)": "F8CBAD",
    "Solo en CE (falta en BC)": "F4B183",
    "Solo en BC (no está en el export CE)": "FFC7CE",
    "Sin datos CE (anterior al export)": "E7E6E6",
}
EXPLICACION = {
    "Coincide": "Está en el export CE y BC creó el pedido con el mismo importe.",
    "Coincide con diferencia de importe": "Está en ambos, pero el total BC no coincide con el total CE.",
    "Devolución con pedido BC": "CE lo da como devuelto (total o parcialmente) y BC había creado el pedido: comprobar abono/devolución en BC.",
    "Devolución omitida en BC (sin pedido)": "CE lo da como devuelto y BC lo omitió (ya estaba RETURNED al procesarlo): no hay pedido BC ni nada que abonar.",
    "Omitido en BC (sin pedido)": "BC lo omitió por su estado CE (CLOSED / MANCO): no hay pedido BC.",
    "Error BC (sin pedido BC)": "Está en CE pero BC no pudo crear el pedido (error de proceso): falta en BC.",
    "Solo en CE (falta en BC)": "Está en el export CE y no aparece en el log de BC.",
    "Solo en BC (no está en el export CE)": "Está en el log de BC con fecha dentro de la cobertura del export, pero el export CE no lo trae.",
    "Sin datos CE (anterior al export)": "Pedido BC anterior al primer pedido del export CE: no se puede comprobar con este export.",
}
COLOR_DEV = {"Sí": "C9A9E0", "Parcial": "C9A9E0", "Posible": "E4DFEC"}
COLOR_SIT = {
    "Pedido BC creado: comprobar abono/devolución en BC": "D9C3EA",
    "Omitido en BC: sin pedido, nada que abonar": "DDEBF7",
    "Error BC: sin pedido en BC": "F8CBAD",
    "Sin pedido en BC": "F4B183",
    "Posible devolución parcial (CLOSED): comprobar en CE": "E4DFEC",
}

CAB_FILL = PatternFill("solid", fgColor="1F4E78")
CAB_FONT = Font(name=FUENTE, size=10, bold=True, color="FFFFFF")
NORMAL = Font(name=FUENTE, size=10)
BOLD = Font(name=FUENTE, size=10, bold=True)
TITULO = Font(name=FUENTE, size=14, bold=True, color="1F4E78")
NOTA = Font(name=FUENTE, size=9, italic=True, color="595959")
AZUL = Font(name=FUENTE, size=10, color="0000FF")
THIN = Side(style="thin", color="BFBFBF")
BORDE = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FMT_FECHA = "dd/mm/yyyy hh:mm"
FMT_DIA = "dd/mm/yyyy"
FMT_EUR = '#,##0.00 "€"'
FMT_DIF = '#,##0.00 "€";[Red]-#,##0.00 "€";"-"'
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
         "septiembre", "octubre", "noviembre", "diciembre"]


# ----------------------------------------------------------------------------- lectura
def leer_log_bc(ruta: Path) -> pd.DataFrame:
    wb = openpyxl.load_workbook(ruta, data_only=False)
    ws = wb["Cruce pedidos"] if "Cruce pedidos" in wb.sheetnames else wb.worksheets[0]
    fila_cab = None
    for r in ws.iter_rows(min_row=1, max_row=30):
        vals = [c.value for c in r]
        if "Channel Order No" in vals:
            fila_cab = r[0].row
            break
    if fila_cab is None:
        sys.exit("No encuentro la cabecera 'Channel Order No' en el Excel de BC")
    cab = [c.value for c in ws[fila_cab]]
    filas = [list(r) for r in ws.iter_rows(min_row=fila_cab + 1, max_row=ws.max_row, values_only=True)]
    df = pd.DataFrame(filas, columns=cab)
    df = df[df["Channel Order No"].notna() & (df["Channel Order No"].astype(str).str.strip() != "")].copy()
    # solo columnas de datos (las de cruce del propio template dependen de su hoja 'Export CE')
    quitar = [c for c in ["Clave cruce", "En export CE", "Resultado BC", "Diferencia importe"] if c in df.columns]
    df = df.drop(columns=quitar)
    for c in ["Merchant Order No", "Channel Order No", "Nº pedido BC", "Error BC"]:
        if c in df.columns:
            df[c] = df[c].map(lambda v: "" if v is None or (isinstance(v, float) and pd.isna(v)) else str(v).strip())
    df["Total BC"] = pd.to_numeric(df["Total BC"], errors="coerce").fillna(0.0)
    no_vacios = df.loc[df["Nº pedido BC"] != "", "Nº pedido BC"]
    if len(no_vacios) and no_vacios.str.fullmatch(r"\d+").all():
        df["Nº pedido BC"] = df["Nº pedido BC"].map(lambda v: int(v) if v else None)
    df["Fecha pedido"] = pd.to_datetime(df["Fecha pedido"], errors="coerce")
    if df["Channel Order No"].duplicated().any():
        dup = df.loc[df["Channel Order No"].duplicated(), "Channel Order No"].tolist()
        print(f"AVISO: Channel Order No duplicados en el log BC: {dup}")
    return df.reset_index(drop=True)


def _fecha_ce_utc(s: pd.Series) -> pd.Series:
    """'09/22/2026 14:35:14 +02:00' -> datetime naive en UTC."""
    m = s.str.extract(r"^(\d\d/\d\d/\d{4} \d\d:\d\d:\d\d) ([+-])(\d\d):(\d\d)$")
    base = pd.to_datetime(m[0], format="%m/%d/%Y %H:%M:%S", errors="coerce")
    signo = m[1].map({"+": 1, "-": -1}).fillna(0).astype(float)
    minutos = pd.to_numeric(m[2], errors="coerce").fillna(0) * 60 + pd.to_numeric(m[3], errors="coerce").fillna(0)
    return base - pd.to_timedelta(minutos * signo, unit="m")


def leer_export_ce(ruta: Path):
    with open(ruta, "r", encoding="utf-8-sig") as f:
        primera = f.readline().strip()
    skip = 1 if primera.lower().startswith("sep=") else 0
    sep = primera[4:5] if skip else ";"
    lineas = pd.read_csv(ruta, sep=sep, skiprows=skip, dtype=str, keep_default_na=False, encoding="utf-8-sig")
    for c in ["Line.Quantity", "Line.LineTotalInclVat", "Line.LineTotalExclVat", "Line.UnitPriceInclVat",
              "Order.ShippingCostsInclVat", "Line.RefundAmountInclVat"]:
        lineas[c] = pd.to_numeric(lineas[c], errors="coerce").fillna(0.0)
    lineas["Fecha pedido (UTC)"] = _fecha_ce_utc(lineas["Order.OrderDate"])
    lineas["Fecha creación CE (UTC)"] = _fecha_ce_utc(lineas["Order.CreatedAt"])
    lineas["Fecha cierre (UTC)"] = _fecha_ce_utc(lineas["Order.ClosedDate"])
    lineas["Fecha envío (UTC)"] = _fecha_ce_utc(lineas["Shipment.ShipmentDate"])
    lineas["Devuelta"] = lineas["Return.Id"].str.strip() != ""
    lineas["Return.Comment"] = lineas["Return.Comment"].str.replace(r"\s*\n\s*>?\s*", " ", regex=True).str.strip(" >")

    def detalle(g):
        partes = []
        for _, l in g[g["Devuelta"]].iterrows():
            partes.append(f"{l['Line.MerchantProductNo']} x{int(l['Line.Quantity'])} {l['Line.LineTotalInclVat']:.2f} € "
                          f"({l['Return.ReturnNo']}, {l['Return.Reason']})")
        return " | ".join(partes)

    def uniq(s):
        return " | ".join(sorted(set(x.strip() for x in s if x and x.strip())))

    grupos = lineas.groupby("Order.ChannelOrderNo", sort=False)
    pedidos = pd.DataFrame({
        "Channel Order No": grupos["Order.ChannelOrderNo"].first(),
        "Merchant Order No": grupos["Order.MerchantOrderNo"].first(),
        "Id CE": grupos["Order.Id"].first(),
        "Marketplace": grupos["Order.ChannelName"].first(),
        "Estado CE": grupos["Order.Status"].agg(uniq),
        "Fecha pedido CE (original)": grupos["Order.OrderDate"].first(),
        "Fecha pedido (UTC)": grupos["Fecha pedido (UTC)"].first(),
        "Fecha creación CE (UTC)": grupos["Fecha creación CE (UTC)"].first(),
        "Fecha cierre (UTC)": grupos["Fecha cierre (UTC)"].first(),
        "Nº líneas": grupos.size(),
        "Unidades": grupos["Line.Quantity"].sum(),
        "Total líneas incl. IVA": grupos["Line.LineTotalInclVat"].sum().round(2),
        "Gastos envío incl. IVA": grupos["Order.ShippingCostsInclVat"].first(),
        "Moneda": grupos["Order.CurrencyCode"].first(),
        "País envío": grupos["ShippingAddress.CountryIso"].first(),
        "País destino AY": grupos["Line.ExtraData.AboutYou.DestinationLocation.CountryCode"].agg(uniq),
        "Líneas devueltas": grupos["Devuelta"].sum().astype(int),
        "Importe líneas devueltas": grupos.apply(lambda g: g.loc[g["Devuelta"], "Line.LineTotalInclVat"].sum(), include_groups=False).round(2),
        "Nº devolución CE": grupos["Return.ReturnNo"].agg(uniq),
        "Motivo devolución": grupos["Return.Reason"].agg(uniq),
        "Comentario devolución": grupos["Return.Comment"].agg(uniq),
        "Detalle líneas devueltas": grupos.apply(detalle, include_groups=False),
    }).reset_index(drop=True)
    pedidos["Total pedido"] = (pedidos["Total líneas incl. IVA"] + pedidos["Gastos envío incl. IVA"]).round(2)
    pedidos = pedidos.sort_values("Fecha pedido (UTC)", ascending=False).reset_index(drop=True)
    return pedidos, lineas


# ----------------------------------------------------------------------------- escritura
def _cabecera(ws, fila, cabeceras, anchos=None):
    for j, h in enumerate(cabeceras, start=1):
        c = ws.cell(row=fila, column=j, value=h)
        c.fill, c.font, c.border = CAB_FILL, CAB_FONT, BORDE
        c.alignment = Alignment(vertical="center", wrap_text=True)
        if anchos and j - 1 < len(anchos) and anchos[j - 1]:
            ws.column_dimensions[get_column_letter(j)].width = anchos[j - 1]
    ws.row_dimensions[fila].height = 30


def _celda(ws, fila, col, valor, fmt=None, font=NORMAL, borde=True):
    c = ws.cell(row=fila, column=col, value=valor)
    c.font = font
    if borde:
        c.border = BORDE
    if fmt:
        c.number_format = fmt
    return c


def _pinta_filas(ws, col_clave, primera, ultima, ncols, colores):
    """Formato condicional: relleno de toda la fila según el texto de col_clave."""
    rango = f"A{primera}:{get_column_letter(ncols)}{ultima}"
    letra = get_column_letter(col_clave)
    for texto, rgb in colores.items():
        ws.conditional_formatting.add(
            rango, FormulaRule(formula=[f'${letra}{primera}="{texto}"'], fill=PatternFill("solid", fgColor=rgb, bgColor=rgb)))


def _escribe_df(ws, df, fila_cab, anchos, formatos):
    _cabecera(ws, fila_cab, list(df.columns), anchos)
    for i, fila in enumerate(df.itertuples(index=False), start=fila_cab + 1):
        for j, v in enumerate(fila, start=1):
            if isinstance(v, float) and pd.isna(v):
                v = None
            elif isinstance(v, pd.Timestamp):
                v = None if pd.isna(v) else v.to_pydatetime()
            elif v is pd.NaT:
                v = None
            _celda(ws, i, j, v, formatos.get(df.columns[j - 1]))
    ws.freeze_panes = ws.cell(row=fila_cab + 1, column=2)
    ws.auto_filter.ref = f"A{fila_cab}:{get_column_letter(len(df.columns))}{max(fila_cab + 1, fila_cab + len(df))}"


def _tabla_cabecera(ws, r, textos):
    _cabecera(ws, r, textos)
    ws.cell(row=r, column=1).fill = PatternFill(fill_type=None)
    ws.cell(row=r, column=1).border = Border()


def _sin_avisos_numero_texto(ruta: Path):
    """Añade <ignoredErrors> a cada hoja para que Excel no marque los identificadores numéricos guardados como texto
    (Merchant Order No de Bol.com). openpyxl 3.1 no expone este elemento, así que se parchea el XML del zip."""
    tag = f'<ignoredErrors><ignoredError sqref="A1:AZ{MAXF}" numberStoredAsText="1"/></ignoredErrors>'
    tmp = ruta.with_suffix(".tmp.xlsx")
    with zipfile.ZipFile(ruta) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            datos = zin.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet") and item.filename.endswith(".xml"):
                xml = datos.decode("utf-8")
                if "<ignoredErrors" not in xml and "</worksheet>" in xml and "<tableParts" not in xml and "<drawing" not in xml:
                    datos = xml.replace("</worksheet>", tag + "</worksheet>").encode("utf-8")
            zout.writestr(item, datos)
    shutil.move(tmp, ruta)


def construir(bc, pedidos, lineas, periodo, salida, ruta_bc, ruta_ce):
    anyo, mes = (int(x) for x in periodo.split("-"))
    ini = dt.datetime(anyo, mes, 1)
    fin = dt.datetime(anyo, mes, calendar.monthrange(anyo, mes)[1], 23, 59, 59)
    nombre_mes = MESES[mes - 1]

    ce_periodo = pedidos[(pedidos["Fecha pedido (UTC)"] >= ini) & (pedidos["Fecha pedido (UTC)"] <= fin)]
    ce_fuera = len(pedidos) - len(ce_periodo)

    # --- unión de claves (log BC + pedidos CE del periodo) -------------------------
    claves_bc = dict(zip(bc["Channel Order No"], zip(bc["Merchant Order No"], bc["Marketplace"], bc["Fecha pedido"])))
    claves_ce = dict(zip(ce_periodo["Channel Order No"],
                         zip(ce_periodo["Merchant Order No"], ce_periodo["Marketplace"], ce_periodo["Fecha pedido (UTC)"])))
    union = []
    for k, (mo, mk, f) in claves_bc.items():
        union.append((k, mo or claves_ce.get(k, ("", "", None))[0], mk, f))
    for k, (mo, mk, f) in claves_ce.items():
        if k not in claves_bc:
            union.append((k, mo, mk, f))
    union.sort(key=lambda t: (t[3] is not None, t[3] if t[3] is not None else dt.datetime.min), reverse=True)

    wb = openpyxl.Workbook()
    ws_res = wb.active
    ws_res.title = "Resumen"
    ws_cru = wb.create_sheet("Cruce pedidos")
    ws_dev = wb.create_sheet("Devoluciones")
    ws_ce = wb.create_sheet("Pedidos CE")
    ws_lin = wb.create_sheet("Líneas CE")
    ws_bc = wb.create_sheet("Log BC")

    # --- hoja Pedidos CE ------------------------------------------------------------
    cols_ce = ["Channel Order No", "Merchant Order No", "Id CE", "Marketplace", "Estado CE", "Fecha pedido CE (original)",
               "Fecha pedido (UTC)", "Fecha creación CE (UTC)", "Fecha cierre (UTC)", "Nº líneas", "Unidades",
               "Total líneas incl. IVA", "Gastos envío incl. IVA", "Total pedido", "Moneda", "País envío", "País destino AY",
               "Líneas devueltas", "Importe líneas devueltas", "Nº devolución CE", "Motivo devolución", "Comentario devolución",
               "Detalle líneas devueltas"]
    pce = pedidos[cols_ce].copy()
    pce["En periodo"] = None  # fórmula
    ws_ce["A1"] = f"Pedidos del export de ChannelEngine (una fila por pedido, agregado de las líneas de {ruta_ce.name})"
    ws_ce["A1"].font = TITULO
    ws_ce["A2"] = ("Fechas convertidas a UTC (About You exporta en +00:00 y Bol.com en +02:00); "
                   "el log de BC guarda la fecha de pedido en UTC.")
    ws_ce["A2"].font = NOTA
    FC = 4
    _escribe_df(ws_ce, pce, FC,
                [22, 18, 9, 26, 12, 24, 17, 17, 17, 8, 9, 14, 12, 14, 8, 8, 10, 9, 13, 30, 16, 40, 60, 10],
                {"Fecha pedido (UTC)": FMT_FECHA, "Fecha creación CE (UTC)": FMT_FECHA, "Fecha cierre (UTC)": FMT_FECHA,
                 "Total líneas incl. IVA": FMT_EUR, "Gastos envío incl. IVA": FMT_EUR, "Total pedido": FMT_EUR,
                 "Importe líneas devueltas": FMT_EUR})
    CE = {c: get_column_letter(i + 1) for i, c in enumerate(pce.columns)}
    for r in range(FC + 1, FC + 1 + len(pce)):
        c = ws_ce[f"{CE['En periodo']}{r}"]
        c.value = (f'=IF(AND({CE["Fecha pedido (UTC)"]}{r}>=Resumen!$C$5,'
                   f'{CE["Fecha pedido (UTC)"]}{r}<=Resumen!$C$6),"Sí","No")')
        c.font, c.border = NORMAL, BORDE

    def rce(col):
        return f"'Pedidos CE'!${CE[col]}${FC + 1}:${CE[col]}${MAXF}"

    # --- hoja Líneas CE (sin datos personales) -------------------------------------
    cols_lin = ["Order.ChannelOrderNo", "Order.MerchantOrderNo", "Order.Id", "Order.ChannelName", "Order.Status",
                "Fecha pedido (UTC)", "Fecha creación CE (UTC)", "Fecha cierre (UTC)", "Fecha envío (UTC)",
                "Line.MerchantProductNo", "Line.Ean", "Line.ProductName", "Line.Quantity", "Line.UnitPriceInclVat",
                "Line.LineTotalExclVat", "Line.LineTotalInclVat", "Line.VatRate", "Line.RefundAmountInclVat",
                "Return.Id", "Return.ReturnNo", "Return.Reason", "Return.Comment", "Return.Origin",
                "ShippingAddress.CountryIso", "Line.ExtraData.AboutYou.SourceLocation.CountryCode",
                "Line.ExtraData.AboutYou.DestinationLocation.CountryCode", "Order.InvoiceNo", "Order.CommercialOrderNo"]
    lin = lineas[cols_lin].copy()
    ws_lin["A1"] = f"Líneas del export de ChannelEngine ({ruta_ce.name}) — solo columnas de negocio, sin datos personales del cliente"
    ws_lin["A1"].font = TITULO
    _escribe_df(ws_lin, lin, 3,
                [22, 16, 8, 24, 12, 17, 17, 17, 17, 22, 15, 30, 6, 10, 10, 10, 7, 9, 8, 22, 12, 40, 9, 8, 8, 8, 14, 16],
                {"Fecha pedido (UTC)": FMT_FECHA, "Fecha creación CE (UTC)": FMT_FECHA, "Fecha cierre (UTC)": FMT_FECHA,
                 "Fecha envío (UTC)": FMT_FECHA, "Line.UnitPriceInclVat": FMT_EUR, "Line.LineTotalExclVat": FMT_EUR,
                 "Line.LineTotalInclVat": FMT_EUR, "Line.RefundAmountInclVat": FMT_EUR})

    # --- hoja Log BC ----------------------------------------------------------------
    ws_bc["A1"] = f"Log de la integración ChannelEngine → Business Central ({ruta_bc.name}, hoja 'Cruce pedidos'), valores tal cual"
    ws_bc["A1"].font = TITULO
    _escribe_df(ws_bc, bc, 3, [18, 22, 26, 12, 17, 12, 18, 10, 10, 10, 16, 60, 14],
                {"Fecha pedido": FMT_FECHA, "Total BC": FMT_EUR})
    BC = {c: get_column_letter(i + 1) for i, c in enumerate(bc.columns)}

    def rbc(col):
        return f"'Log BC'!${BC[col]}$4:${BC[col]}${MAXF}"

    # --- hoja Cruce pedidos ---------------------------------------------------------
    ws_cru["A1"] = f"Cruce de pedidos ChannelEngine vs Business Central — {nombre_mes} {anyo}"
    ws_cru["A1"].font = TITULO
    ws_cru["A2"] = ("Una fila por pedido (unión del log BC y de los pedidos del export CE con fecha de pedido dentro del periodo). "
                    "Las columnas E en adelante son fórmulas sobre las hojas 'Pedidos CE' y 'Log BC'; el color de la fila sigue "
                    "la columna Resultado y la columna Devolución se marca en lila.")
    ws_cru["A2"].font = NOTA
    cab = ["Channel Order No", "Merchant Order No", "Marketplace", "Fecha pedido (UTC)",
           "En export CE", "Estado CE (export)", "Total CE",
           "En log BC", "Estado CE (log BC)", "Total BC", "Dif. importe (BC − CE)",
           "Estado proceso BC", "Nº pedido BC", "Error BC", "Pedido BC creado",
           "Devolución", "Detalle devolución (export CE)", "Resultado"]
    FCR = 4
    _cabecera(ws_cru, FCR, cab, [22, 16, 24, 17, 9, 12, 12, 8, 12, 12, 13, 13, 12, 44, 10, 10, 50, 36])
    L = {c: get_column_letter(i + 1) for i, c in enumerate(cab)}
    p1 = FCR + 1
    for i, (k, mo, mk, f) in enumerate(union, start=p1):
        _celda(ws_cru, i, 1, k)
        _celda(ws_cru, i, 2, mo or None)
        _celda(ws_cru, i, 3, mk)
        _celda(ws_cru, i, 4, f.to_pydatetime() if isinstance(f, pd.Timestamp) else f, FMT_FECHA)
        A = f"$A{i}"
        e, h, l_ = f'{L["En export CE"]}{i}', f'{L["En log BC"]}{i}', f'{L["Estado proceso BC"]}{i}'
        formulas = {
            "En export CE": f'=IF(COUNTIF({rce("Channel Order No")},{A})>0,"Sí","No")',
            "Estado CE (export)": f'=IF({e}="Sí",INDEX({rce("Estado CE")},MATCH({A},{rce("Channel Order No")},0))&"","")',
            "Total CE": f'=IF({e}="Sí",SUMIF({rce("Channel Order No")},{A},{rce("Total pedido")}),"")',
            "En log BC": f'=IF(COUNTIF({rbc("Channel Order No")},{A})>0,"Sí","No")',
            "Estado CE (log BC)": f'=IF({h}="Sí",INDEX({rbc("Estado CE")},MATCH({A},{rbc("Channel Order No")},0))&"","")',
            "Total BC": f'=IF({h}="Sí",SUMIF({rbc("Channel Order No")},{A},{rbc("Total BC")}),"")',
            "Dif. importe (BC − CE)": f'=IF(AND({e}="Sí",{h}="Sí"),ROUND({L["Total BC"]}{i}-{L["Total CE"]}{i},2),"")',
            "Estado proceso BC": f'=IF({h}="Sí",INDEX({rbc("Estado proceso BC")},MATCH({A},{rbc("Channel Order No")},0))&"","")',
            "Nº pedido BC": (f'=IF({h}="Sí",IF(INDEX({rbc("Nº pedido BC")},MATCH({A},{rbc("Channel Order No")},0))="","",'
                             f'INDEX({rbc("Nº pedido BC")},MATCH({A},{rbc("Channel Order No")},0))),"")'),
            "Error BC": f'=IF({h}="Sí",INDEX({rbc("Error BC")},MATCH({A},{rbc("Channel Order No")},0))&"","")',
            "Pedido BC creado": f'=IF(AND({l_}="Processed",{L["Nº pedido BC"]}{i}<>""),"Sí","No")',
            "Detalle devolución (export CE)": (f'=IF({e}="Sí",INDEX({rce("Detalle líneas devueltas")},'
                                              f'MATCH({A},{rce("Channel Order No")},0))&"","")'),
            "Devolución": (f'=IF(OR({L["Estado CE (export)"]}{i}="RETURNED",AND({e}="No",{L["Estado CE (log BC)"]}{i}="RETURNED")),"Sí",'
                           f'IF({L["Detalle devolución (export CE)"]}{i}<>"","Parcial",'
                           f'IF(AND({e}="No",{L["Estado CE (log BC)"]}{i}="CLOSED"),"Posible","No")))'),
            "Resultado": (f'=IF(AND({h}="No",{e}="Sí"),"Solo en CE (falta en BC)",'
                          f'IF({e}="No",IF({L["Fecha pedido (UTC)"]}{i}<Resumen!$C$8,"Sin datos CE (anterior al export)",'
                          f'"Solo en BC (no está en el export CE)"),'
                          f'IF({l_}="Error","Error BC (sin pedido BC)",'
                          f'IF({l_}="Skipped",IF({L["Devolución"]}{i}<>"No","Devolución omitida en BC (sin pedido)",'
                          f'"Omitido en BC (sin pedido)"),'
                          f'IF({L["Devolución"]}{i}<>"No","Devolución con pedido BC",'
                          f'IF(AND({L["Dif. importe (BC − CE)"]}{i}<>"",{L["Dif. importe (BC − CE)"]}{i}<>0),'
                          f'"Coincide con diferencia de importe","Coincide"))))))'),
        }
        fmts = {"Total CE": FMT_EUR, "Total BC": FMT_EUR, "Dif. importe (BC − CE)": FMT_DIF}
        for col, fx in formulas.items():
            _celda(ws_cru, i, cab.index(col) + 1, fx, fmts.get(col))
    pN = p1 + len(union) - 1
    ws_cru.freeze_panes = f"B{p1}"
    ws_cru.auto_filter.ref = f"A{FCR}:{get_column_letter(len(cab))}{pN}"
    # colores: primero la columna Devolución (prioridad alta), luego la fila por Resultado
    cD = L["Devolución"]
    for texto, rgb in COLOR_DEV.items():
        ws_cru.conditional_formatting.add(
            f"{cD}{p1}:{cD}{pN}",
            FormulaRule(formula=[f'${cD}{p1}="{texto}"'], fill=PatternFill("solid", fgColor=rgb, bgColor=rgb),
                        font=Font(name=FUENTE, size=10, bold=True)))
    _pinta_filas(ws_cru, cab.index("Resultado") + 1, p1, pN, len(cab), COLOR)

    def rcr(col):
        return f"'Cruce pedidos'!${L[col]}${p1}:${L[col]}${pN}"

    # --- hoja Devoluciones ----------------------------------------------------------
    ws_dev["A1"] = f"Devoluciones — {nombre_mes} {anyo}"
    ws_dev["A1"].font = TITULO
    ws_dev["A2"] = ("Pedidos del cruce con Devolución = Sí / Parcial / Posible. 'Origen del dato' indica si la devolución está "
                    "confirmada por el export CE o solo consta en el log de BC. El color de la fila sigue 'Situación'.")
    ws_dev["A2"].font = NOTA
    cab_d = ["Channel Order No", "Merchant Order No", "Marketplace", "Fecha pedido (UTC)", "Origen del dato", "Devolución",
             "Estado CE (export)", "Estado CE (log BC)", "Total pedido", "Líneas devueltas", "Importe líneas devueltas",
             "Nº devolución CE", "Motivo", "Detalle líneas devueltas", "Estado proceso BC", "Nº pedido BC", "Situación"]
    FCD = 4
    _cabecera(ws_dev, FCD, cab_d, [22, 16, 24, 17, 13, 10, 12, 12, 12, 9, 12, 30, 14, 55, 13, 12, 50])
    D = {c: get_column_letter(i + 1) for i, c in enumerate(cab_d)}
    # selección de filas en Python con la misma regla que la fórmula 'Devolución'
    idx_ce = ce_periodo.set_index("Channel Order No")
    idx_bc = bc.set_index("Channel Order No")

    def es_devolucion(k):
        en_ce = k in idx_ce.index
        est_ce = idx_ce.loc[k, "Estado CE"] if en_ce else ""
        est_bc = idx_bc.loc[k, "Estado CE"] if k in idx_bc.index else ""
        det = idx_ce.loc[k, "Detalle líneas devueltas"] if en_ce else ""
        if est_ce == "RETURNED" or (not en_ce and est_bc == "RETURNED"):
            return True
        if det:
            return True
        return (not en_ce) and est_bc == "CLOSED"

    devs = [u for u in union if es_devolucion(u[0])]
    d1 = FCD + 1
    for i, (k, mo, mk, f) in enumerate(devs, start=d1):
        _celda(ws_dev, i, 1, k)
        _celda(ws_dev, i, 2, mo or None)
        _celda(ws_dev, i, 3, mk)
        _celda(ws_dev, i, 4, f.to_pydatetime() if isinstance(f, pd.Timestamp) else f, FMT_FECHA)
        A = f"$A{i}"

        def m_cru(col):
            return f'INDEX({rcr(col)},MATCH({A},{rcr("Channel Order No")},0))'

        def m_ce(col):
            return f'INDEX({rce(col)},MATCH({A},{rce("Channel Order No")},0))'

        o = f'{D["Origen del dato"]}{i}'
        fx = {
            "Origen del dato": f'=IF({m_cru("En export CE")}="Sí","Export CE","Solo log BC")',
            "Devolución": f'={m_cru("Devolución")}&""',
            "Estado CE (export)": f'={m_cru("Estado CE (export)")}&""',
            "Estado CE (log BC)": f'={m_cru("Estado CE (log BC)")}&""',
            "Total pedido": f'=IF({m_cru("En log BC")}="Sí",{m_cru("Total BC")},{m_cru("Total CE")})',
            "Líneas devueltas": f'=IF({o}="Export CE",{m_ce("Líneas devueltas")},"")',
            "Importe líneas devueltas": f'=IF({o}="Export CE",{m_ce("Importe líneas devueltas")},"")',
            "Nº devolución CE": f'=IF({o}="Export CE",{m_ce("Nº devolución CE")}&"","")',
            "Motivo": f'=IF({o}="Export CE",{m_ce("Motivo devolución")}&"","")',
            "Detalle líneas devueltas": f'=IF({o}="Export CE",{m_ce("Detalle líneas devueltas")}&"","")',
            "Estado proceso BC": f'={m_cru("Estado proceso BC")}&""',
            "Nº pedido BC": f'=IF({m_cru("Nº pedido BC")}="","",{m_cru("Nº pedido BC")})',
            "Situación": (f'=IF({D["Devolución"]}{i}="Posible","Posible devolución parcial (CLOSED): comprobar en CE",'
                          f'IF({D["Nº pedido BC"]}{i}<>"","Pedido BC creado: comprobar abono/devolución en BC",'
                          f'IF({D["Estado proceso BC"]}{i}="Skipped","Omitido en BC: sin pedido, nada que abonar",'
                          f'IF({D["Estado proceso BC"]}{i}="Error","Error BC: sin pedido en BC","Sin pedido en BC"))))'),
        }
        fmts = {"Total pedido": FMT_EUR, "Importe líneas devueltas": FMT_EUR}
        for col, f_ in fx.items():
            _celda(ws_dev, i, cab_d.index(col) + 1, f_, fmts.get(col))
    dN = d1 + max(len(devs), 1) - 1
    ws_dev.freeze_panes = f"B{d1}"
    ws_dev.auto_filter.ref = f"A{FCD}:{get_column_letter(len(cab_d))}{dN}"
    cO = D["Origen del dato"]
    ws_dev.conditional_formatting.add(
        f"{cO}{d1}:{cO}{dN}",
        FormulaRule(formula=[f'${cO}{d1}="Solo log BC"'], fill=PatternFill("solid", fgColor="E7E6E6", bgColor="E7E6E6"),
                    font=Font(name=FUENTE, size=10, italic=True)))
    _pinta_filas(ws_dev, cab_d.index("Situación") + 1, d1, dN, len(cab_d), COLOR_SIT)

    def rdv(col):
        return f"'Devoluciones'!${D[col]}${d1}:${D[col]}${dN}"

    # --- hoja Resumen ---------------------------------------------------------------
    ws = ws_res
    for col, w in zip("ABCDEF", [3, 60, 18, 18, 18, 95]):
        ws.column_dimensions[col].width = w
    ws["B1"] = f"Cuadre ChannelEngine vs Business Central — {nombre_mes} {anyo}"
    ws["B1"].font = TITULO
    ws["B2"] = (f"Generado el {dt.date.today():%d/%m/%Y} a partir de '{ruta_bc.name}' (log BC, hoja 'Cruce pedidos') "
                f"y '{ruta_ce.name}' (export CE). Celdas en azul: parámetros editables.")
    ws["B2"].font = NOTA
    ws["B4"] = "Parámetros"
    ws["B4"].font = BOLD
    filas_param = [
        ("Periodo: inicio", ini, FMT_DIA, AZUL, "Fecha de pedido (UTC) desde la que se cruza. Editable."),
        ("Periodo: fin", fin, FMT_FECHA, AZUL, "Fecha de pedido (UTC) hasta la que se cruza. Editable."),
        ("Pedidos en el export CE (total)", f"=COUNTA({rce('Channel Order No')})", "0", NORMAL,
         "Todos los pedidos del CSV, dentro y fuera del periodo."),
        ("Export CE: primer pedido (fecha pedido UTC)", f"=MIN({rce('Fecha pedido (UTC)')})", FMT_FECHA, NORMAL,
         "Inicio de la cobertura del export. Los pedidos BC anteriores se marcan 'Sin datos CE'. Se puede sobrescribir a mano."),
        ("Export CE: último pedido (fecha pedido UTC)", f"=MAX({rce('Fecha pedido (UTC)')})", FMT_FECHA, NORMAL, ""),
        ("Export CE: primera fecha de creación en CE", f"=MIN({rce('Fecha creación CE (UTC)')})", FMT_FECHA, NORMAL,
         "Si es posterior al inicio del periodo, el export se descargó filtrado por fecha de creación y no cubre todo el mes."),
        ("Pedidos del export fuera del periodo (no se cruzan)", f'=COUNTIF({rce("En periodo")},"No")', "0", NORMAL,
         "Quedan en la hoja 'Pedidos CE' pero no entran en el cruce."),
    ]
    r = 5
    for etiqueta, valor, fmt, fnt, nota in filas_param:
        _celda(ws, r, 2, etiqueta)
        _celda(ws, r, 3, valor, fmt, fnt)
        _celda(ws, r, 6, nota, font=NOTA, borde=False)
        r += 1
    # C5 inicio periodo, C6 fin, C7 total export, C8 primer pedido export, C9 último, C10 creación, C11 fuera periodo
    # (las fórmulas de arriba referencian Resumen!$C$5 / $C$6 / $C$8: se comprueba aquí)
    assert ws["B6"].value == "Periodo: fin" and ws["B5"].value == "Periodo: inicio"
    assert ws["B8"].value.startswith("Export CE: primer pedido")
    def _fecha(c):  # TEXT(fecha,"dd/mm/yyyy") depende del idioma de Excel; DAY/MONTH/YEAR no
        return f'TEXT(DAY({c}),"00")&"/"&TEXT(MONTH({c}),"00")&"/"&YEAR({c})'

    ws["B13"] = (f'=IF(C8>C5,"AVISO: el export CE empieza el "&{_fecha("C8")}&" "&TEXT(HOUR(C8),"00")&":"&TEXT(MINUTE(C8),"00")'
                 f'&". Los pedidos BC anteriores no se pueden comprobar: vuelve a descargar el informe de pedidos de CE '
                 f'filtrando por fecha de pedido del "&{_fecha("C5")}&" al "&{_fecha("C6")}&".",'
                 f'"El export CE cubre todo el periodo.")')
    ws["B13"].font = Font(name=FUENTE, size=10, bold=True, color="C00000")
    ws.merge_cells("B13:F13")
    ws["B13"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[13].height = 32

    r = 15
    ws.cell(row=r, column=2, value="Pedidos").font = BOLD
    r += 1
    _tabla_cabecera(ws, r, ["", "Concepto", "Pedidos", "Importe BC", "Importe CE", "Cómo se calcula"])
    r += 1
    ambos_bc = f'SUMIFS({rcr("Total BC")},{rcr("En export CE")},"Sí",{rcr("En log BC")},"Sí")'
    ambos_ce = f'SUMIFS({rcr("Total CE")},{rcr("En export CE")},"Sí",{rcr("En log BC")},"Sí")'
    kpis = [
        ("Pedidos en el log BC (periodo)", f'=COUNTIF({rcr("En log BC")},"Sí")',
         f'=SUMIF({rcr("En log BC")},"Sí",{rcr("Total BC")})', None, "Filas del log de BC."),
        ("   de los que BC creó pedido (Processed)", f'=COUNTIF({rcr("Pedido BC creado")},"Sí")',
         f'=SUMIF({rcr("Pedido BC creado")},"Sí",{rcr("Total BC")})', None, "Estado proceso BC = Processed con nº de pedido."),
        ("   con error de proceso en BC (sin pedido)", f'=COUNTIF({rcr("Estado proceso BC")},"Error")',
         f'=SUMIF({rcr("Estado proceso BC")},"Error",{rcr("Total BC")})', None, "Estado proceso BC = Error."),
        ("   omitidos por BC (sin pedido)", f'=COUNTIF({rcr("Estado proceso BC")},"Skipped")',
         f'=SUMIF({rcr("Estado proceso BC")},"Skipped",{rcr("Total BC")})', None,
         "Estado proceso BC = Skipped (RETURNED / CLOSED / MANCO al procesar)."),
        ("Pedidos del export CE dentro del periodo", f'=COUNTIF({rcr("En export CE")},"Sí")', None,
         f'=SUMIF({rcr("En export CE")},"Sí",{rcr("Total CE")})', "Pedidos del CSV con fecha de pedido (UTC) dentro del periodo."),
        ("   en ambos (export CE y log BC)", f'=COUNTIFS({rcr("En export CE")},"Sí",{rcr("En log BC")},"Sí")',
         f"={ambos_bc}", f"={ambos_ce}", ""),
        ("   diferencia de importe en los pedidos en ambos", None, f"={ambos_bc}-{ambos_ce}", None,
         "Importe BC − importe CE de los pedidos presentes en ambos. Debe ser 0."),
    ]
    for etiqueta, n, imp_bc, imp_ce, nota in kpis:
        _celda(ws, r, 2, etiqueta)
        _celda(ws, r, 3, n, "0")
        _celda(ws, r, 4, imp_bc, FMT_EUR)
        _celda(ws, r, 5, imp_ce, FMT_EUR)
        _celda(ws, r, 6, nota, font=NOTA, borde=False)
        r += 1

    r += 1
    ws.cell(row=r, column=2, value="Resultado del cruce (color de fila en 'Cruce pedidos')").font = BOLD
    r += 1
    _tabla_cabecera(ws, r, ["", "Resultado", "Pedidos", "Importe BC", "Importe CE", "Significado"])
    r += 1
    r_ini = r
    for texto, rgb in COLOR.items():
        c = _celda(ws, r, 2, texto)
        c.fill = PatternFill("solid", fgColor=rgb)
        _celda(ws, r, 3, f'=COUNTIF({rcr("Resultado")},"{texto}")', "0")
        _celda(ws, r, 4, f'=SUMIF({rcr("Resultado")},"{texto}",{rcr("Total BC")})', FMT_EUR)
        _celda(ws, r, 5, f'=SUMIF({rcr("Resultado")},"{texto}",{rcr("Total CE")})', FMT_EUR)
        _celda(ws, r, 6, EXPLICACION[texto], font=NOTA, borde=False)
        r += 1
    _celda(ws, r, 2, "Total filas del cruce", font=BOLD)
    _celda(ws, r, 3, f'=COUNTA({rcr("Channel Order No")})', "0", BOLD)
    _celda(ws, r, 4, f"=SUM(D{r_ini}:D{r - 1})", FMT_EUR, BOLD)
    _celda(ws, r, 5, f"=SUM(E{r_ini}:E{r - 1})", FMT_EUR, BOLD)
    r += 2

    ws.cell(row=r, column=2, value="Devoluciones (hoja 'Devoluciones')").font = BOLD
    r += 1
    _tabla_cabecera(ws, r, ["", "Concepto", "Pedidos", "Importe BC", "Importe CE", "Cómo se calcula"])
    r += 1
    s_creado = "Pedido BC creado: comprobar abono/devolución en BC"
    s_omit = "Omitido en BC: sin pedido, nada que abonar"
    devk = [
        ("Pedidos devueltos (Devolución = Sí)", f'=COUNTIF({rcr("Devolución")},"Sí")',
         f'=SUMIF({rcr("Devolución")},"Sí",{rcr("Total BC")})', f'=SUMIF({rcr("Devolución")},"Sí",{rcr("Total CE")})',
         "Estado CE RETURNED (en el export, o en el log BC si el export no lo trae)."),
        ("Pedidos con devolución parcial (CLOSED con líneas devueltas en el export)", f'=COUNTIF({rcr("Devolución")},"Parcial")',
         f'=SUMIF({rcr("Devolución")},"Parcial",{rcr("Total BC")})', f'=SUMIF({rcr("Devolución")},"Parcial",{rcr("Total CE")})',
         "El importe devuelto real está en 'Importe líneas devueltas' de la hoja Devoluciones."),
        ("Pedidos CLOSED solo en el log BC (posible devolución parcial)", f'=COUNTIF({rcr("Devolución")},"Posible")',
         f'=SUMIF({rcr("Devolución")},"Posible",{rcr("Total BC")})', None, "Sin export CE no se sabe si tienen líneas devueltas."),
        ("   devoluciones confirmadas por el export CE", f'=COUNTIF({rdv("Origen del dato")},"Export CE")', None, None, ""),
        ("   devoluciones que solo constan en el log BC", f'=COUNTIF({rdv("Origen del dato")},"Solo log BC")', None, None,
         "Pendientes de confirmar con el export completo."),
        ("   con pedido BC creado → comprobar abono/devolución en BC", f'=COUNTIF({rdv("Situación")},"{s_creado}")',
         f'=SUMIF({rdv("Situación")},"{s_creado}",{rdv("Total pedido")})', None, "Situación = 'Pedido BC creado…'."),
        ("   omitidas por BC (sin pedido, nada que abonar)", f'=COUNTIF({rdv("Situación")},"{s_omit}")',
         f'=SUMIF({rdv("Situación")},"{s_omit}",{rdv("Total pedido")})', None, ""),
        ("   con error BC / sin pedido en BC",
         f'=COUNTIF({rdv("Situación")},"Error BC: sin pedido en BC")+COUNTIF({rdv("Situación")},"Sin pedido en BC")', None, None, ""),
        ("Importe de las líneas devueltas según el export CE", f'=COUNTIF({rdv("Líneas devueltas")},">0")', None,
         f'=SUM({rdv("Importe líneas devueltas")})',
         "Pedidos con líneas devueltas en el export y suma de esas líneas (IVA incluido)."),
    ]
    for etiqueta, n, imp_bc, imp_ce, nota in devk:
        _celda(ws, r, 2, etiqueta)
        _celda(ws, r, 3, n, "0")
        _celda(ws, r, 4, imp_bc, FMT_EUR)
        _celda(ws, r, 5, imp_ce, FMT_EUR)
        _celda(ws, r, 6, nota, font=NOTA, borde=False)
        r += 1
    r += 1
    ws.cell(row=r, column=2, value="Leyenda de la columna Devolución").font = BOLD
    r += 1
    leyenda_dev = {"Sí": "Pedido devuelto (estado CE RETURNED).",
                   "Parcial": "Pedido CLOSED con alguna línea devuelta en el export CE.",
                   "Posible": "Pedido CLOSED solo en el log BC: sin export no se puede saber si hay líneas devueltas."}
    for texto, rgb in COLOR_DEV.items():
        c = _celda(ws, r, 2, texto, font=BOLD)
        c.fill = PatternFill("solid", fgColor=rgb)
        _celda(ws, r, 6, leyenda_dev[texto], font=NOTA, borde=False)
        r += 1
    ws.sheet_view.showGridLines = False

    wb.calculation.fullCalcOnLoad = True
    salida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(salida)
    _sin_avisos_numero_texto(salida)
    return {"filas_cruce": len(union), "devoluciones": len(devs), "pedidos_ce": len(pedidos),
            "pedidos_ce_periodo": len(ce_periodo), "pedidos_ce_fuera": ce_fuera, "log_bc": len(bc)}


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--bc", required=True, help="Excel del log BC (hoja 'Cruce pedidos')")
    ap.add_argument("--ce", required=True, help="CSV orders-report de ChannelEngine")
    ap.add_argument("--periodo", default=None, help="Mes a cuadrar, AAAA-MM (por defecto, el mes del primer pedido del log BC)")
    ap.add_argument("--salida", default=None, help="Excel de salida (por defecto Cuadre_CE_vs_BC_<periodo>.xlsx junto al log BC)")
    a = ap.parse_args()
    ruta_bc, ruta_ce = Path(a.bc), Path(a.ce)
    bc = leer_log_bc(ruta_bc)
    pedidos, lineas = leer_export_ce(ruta_ce)
    periodo = a.periodo or bc["Fecha pedido"].min().strftime("%Y-%m")
    salida = Path(a.salida) if a.salida else ruta_bc.with_name(f"Cuadre_CE_vs_BC_{periodo}.xlsx")
    info = construir(bc, pedidos, lineas, periodo, salida, ruta_bc, ruta_ce)
    print(f"Log BC: {info['log_bc']} pedidos | Export CE: {info['pedidos_ce']} pedidos ({info['pedidos_ce_periodo']} en {periodo}, "
          f"{info['pedidos_ce_fuera']} fuera) | Cruce: {info['filas_cruce']} filas | Devoluciones: {info['devoluciones']}")
    print(f"Guardado: {salida}")


if __name__ == "__main__":
    main()
